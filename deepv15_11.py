# -*- coding: utf-8 -*-
"""
Maçkolik arşivi — v15.11
+ Gün tamamlanma doğrulaması (artık tek maçla gün "tamam" sayılmıyor)
+ Lig imzası değişince checkpoint reddi (eski grup verisi karışmıyor)
+ HTTP kilidi yalnızca throttle kararını kapsıyor (gerçek eşzamanlılık)
+ MatchData sadece başarıda işaretleniyor (transient hatalar tekrar deneniyor)
+ Hata listeleri için yapı doğrulaması
+ process_row tamamen korumalı (gather sızıntısı yok)
"""

import asyncio
import argparse
import os
import re
import json
import time
import random
import threading
import subprocess
import requests
from pathlib import Path
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime

import pandas as pd
import cloudscraper
from bs4 import BeautifulSoup
from curl_cffi.requests import AsyncSession


# ============================================================
# ÜLKE / LİG GRUPLARI
# ============================================================

ULKE_GRUPLARI = {
    "TR": {"ad": "Türkiye", "ligler": {(1, 1), (2, 1)}},
    "EN": {"ad": "İngiltere", "ligler": {(24, 17), (25, 17)}},
    "DE": {"ad": "Almanya", "ligler": {(3, 2), (4, 2)}},
    "FR": {"ad": "Fransa", "ligler": {(5, 3), (6, 3)}},
    "IT": {"ad": "İtalya", "ligler": {(15, 10), (16, 10)}},
    "ES": {"ad": "İspanya", "ligler": {(20, 14), (21, 14)}},
    "PT": {"ad": "Portekiz", "ligler": {(19, 13), (474, 13)}},
    "BE": {"ad": "Belçika", "ligler": {(9, 5), (415, 5)}},
    "NL": {"ad": "Hollanda", "ligler": {(17, 11), (120, 11)}},
    "SC": {"ad": "İskoçya", "ligler": {(26, 18), (106, 18)}},
}

AKTIF_GRUPLAR = []
AKTIF_LIGLER = set()
AKTIF_ETIKET = "ALL"


def _lig_uygun_mu(yarisma):
    if not isinstance(yarisma, list) or len(yarisma) < 3:
        return False
    try:
        ulke_id = int(yarisma[0]); lig_id = int(yarisma[2])
    except (TypeError, ValueError):
        return False
    return (lig_id, ulke_id) in AKTIF_LIGLER


def _lig_filtresi_imzasi():
    parcalar = []
    for kod in sorted(AKTIF_GRUPLAR):
        parcalar.append(kod)
        for lig_id, ulke_id in sorted(ULKE_GRUPLARI[kod]["ligler"]):
            parcalar.append(f"{lig_id}-{ulke_id}")
    return "|".join(parcalar)


# ============================================================
# AYARLAR
# ============================================================

KLASOR = Path(__file__).resolve().parent
ISTEK_ARASI = 10.0
ISTEK_JITTER = 2.5
ISTEK_ARASI_LIVEDATA = 5.0
ISTEK_JITTER_LIVEDATA = 1.5

ZAMAN_ASIMI = 60
EXCEL_KAYIT_ARALIGI = 100
ISTATISTIK_ON_EK = "İSTATİSTİK | "
MAKS_TUR = 3
LOG_AKTIF = True
MATCHDATA_AKTIF = True

# ═══ HTTP THROTTLE ═══
ESZAMANLI_ISTEK = 2
MIN_ISTEK_ARASI = 6.0
ISTEK_JITTER_ASYNC = 2.5

# ═══ WARP OTO-RESET ═══
ARDISIK_429_ESIK = 3
WARP_RESET_SONRASI_BEKLEME = 90
WARP_RESET_MAKS = 3
WARP_PATH_CANDIDATES = [
    "warp-cli",
    r"C:\Program Files\Cloudflare\Cloudflare WARP\warp-cli.exe",
    r"C:\Program Files\Cloudflare\Cloudflare One Client\warp-cli.exe",
]

CURL_IMPERSONATE = "chrome120"
GLOBAL_COOLDOWN_LOG_ESIGI = 5.0

ISTATISTIK_NORMALIZE = {
    "Fauller": "Faul", "Ofsaytlar": "Ofsayt", "Köşe Vuruşu": "Korner",
}

RETRY_AFTER_BILGI_MAKS = 604800
RETRY_AFTER_BEKLEME_MAKS = 900
LIVEDATA_RETRY_AFTER_MAX = 60

DENEME_SAYISI_5XX = 5
DENEME_SAYISI_AG = 4

ISTEK_HEADERS = {
    "Referer": "https://www.mackolik.com/",
    "Accept-Language": "tr-TR,tr;q=0.9,en;q=0.8",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    ),
}

MAC_KOLONLARI = [
    "Tarih", "Saat", "Lig/Sezon", "Ev Sahibi", "Deplasman", "Maç",
    "İY Skor", "İY Skor Kaynağı", "Skor", "Maç Durumu", "Maç ID",
    "İddaa Kodu", "Seyirci", "Stadyum", "Hakem", "TV Yayını",
    "Ev Sahibi Teknik Direktör", "Deplasman Teknik Direktör",
    "Ev Kadro", "Deplasman Kadro", "Ev Yedek", "Deplasman Yedek",
    "Goller", "Sarı Kartlar", "Kırmızı Kartlar", "Değişiklikler",
    "Kaçan Penaltılar",
    "Maç Slug", "Kaynak URL", "Arşiv Sayfa Durumu",
    "Oran Durumu", "İstatistik Durumu"
]

HATA_KOLONLARI = ["Tarih", "ID", "Maç", "Hata"]


# ============================================================
# YARDIMCILAR (erken)
# ============================================================

def metin_temizle(deger):
    if deger is None:
        return ""
    return re.sub(r"\s+", " ", str(deger)).strip()


def sure_formatla(saniye):
    saniye = max(0, int(saniye))
    saat, kalan = divmod(saniye, 3600)
    dakika, saniye = divmod(kalan, 60)
    if saat:
        return f"{saat} sa {dakika} dk {saniye} sn"
    if dakika:
        return f"{dakika} dk {saniye} sn"
    return f"{saniye} sn"


# ============================================================
# HATA DEPOSU
# ============================================================

def hata_anahtari(kayit):
    mid = metin_temizle(kayit.get("ID", ""))
    if mid:
        return f"id:{mid}"
    tarih = metin_temizle(kayit.get("Tarih", ""))
    if tarih:
        return f"tarih:{tarih}"
    return "__genel__"


def hata_ekle(hatalar, kayit):
    """hatalar sözlüğünü YERİNDE mutasyona uğratır."""
    if not isinstance(hatalar, dict):
        return
    kayit_norm = {k: metin_temizle(kayit.get(k, "")) for k in HATA_KOLONLARI}
    anahtar = hata_anahtari(kayit_norm)
    hatalar.setdefault(anahtar, []).append(kayit_norm)


def hata_listesini_dict_yap(hatalar):
    """Eski liste ve bozuk checkpoint yapılarını güvenli sözlüğe dönüştürür."""
    yeni = {}
    if isinstance(hatalar, dict):
        kaynak = []
        for kayitlar in hatalar.values():
            if isinstance(kayitlar, list):
                kaynak.extend(k for k in kayitlar if isinstance(k, dict))
            elif isinstance(kayitlar, dict):
                kaynak.append(kayitlar)
        for kayit in kaynak:
            hata_ekle(yeni, kayit)
    elif isinstance(hatalar, list):
        for kayit in hatalar:
            if isinstance(kayit, dict):
                hata_ekle(yeni, kayit)
    return yeni


def hatalari_birlestir(hedef, kaynak):
    kaynak = hata_listesini_dict_yap(kaynak)
    for anahtar, kayitlar in kaynak.items():
        hedef.setdefault(anahtar, []).extend(kayitlar)
    return hedef


def hatalari_listeye_cevir(hatalar):
    """
    v15.7: Bozuk veri yapısını tolere eder.
    Değer list değilse veya elemanlar dict değilse sessizce atlar.
    """
    if isinstance(hatalar, list):
        return [k for k in hatalar if isinstance(k, dict)]
    sonuc = []
    if not isinstance(hatalar, dict):
        return sonuc
    for kayitlar in hatalar.values():
        if isinstance(kayitlar, list):
            for k in kayitlar:
                if isinstance(k, dict):
                    sonuc.append(k)
        elif isinstance(kayitlar, dict):
            sonuc.append(kayitlar)
    return sonuc


def hata_kayitlarini_temizle(hatalar, mac_id, baslangiclar=(), icerenler=()):
    """
    YERİNDE mutasyon; aynı sözlüğü döndürür.
    """
    if not isinstance(hatalar, dict):
        return hatalar
    anahtar = f"id:{str(mac_id)}"
    kayitlar = hatalar.get(anahtar)
    if not kayitlar:
        return hatalar
    kalan = []
    for kayit in kayitlar:
        if not isinstance(kayit, dict):
            continue
        mesaj = metin_temizle(kayit.get("Hata", ""))
        temizle = any(mesaj.startswith(x) for x in baslangiclar)
        temizle = temizle or any(x.casefold() in mesaj.casefold() for x in icerenler)
        if not temizle:
            kalan.append(kayit)
    if kalan:
        hatalar[anahtar] = kalan
    else:
        hatalar.pop(anahtar, None)
    return hatalar


def hatalari_tekrarsiz_yap(hatalar):
    hatalar = hata_listesini_dict_yap(hatalar)
    for anahtar in list(hatalar.keys()):
        gorulen = set()
        yeni = []
        kayitlar = hatalar.get(anahtar)
        if not isinstance(kayitlar, list):
            kayitlar = [kayitlar] if isinstance(kayitlar, dict) else []
        for kayit in kayitlar:
            if not isinstance(kayit, dict):
                continue
            k = tuple(metin_temizle(kayit.get(f, "")) for f in HATA_KOLONLARI)
            if k not in gorulen:
                gorulen.add(k)
                yeni.append({f: metin_temizle(kayit.get(f, ""))
                             for f in HATA_KOLONLARI})
        if yeni:
            hatalar[anahtar] = yeni
        else:
            hatalar.pop(anahtar, None)
    return hatalar


def tarih_hatalarini_temizle(hatalar, gunler):
    for gun in gunler:
        hatalar.pop(f"tarih:{gun}", None)
    return hatalar


# ============================================================
# DOSYA YOLLARI VE TARİH
# ============================================================

def dosya_yollari_uret(basla, son):
    alt_klasor = KLASOR / f"{basla}_{son}_{AKTIF_ETIKET}_ARSIV_TEK_SAYFA"
    alt_klasor.mkdir(parents=True, exist_ok=True)
    return {
        "klasor": alt_klasor,
        "excel": alt_klasor / "bahis_oranlari.xlsx",
        "hata": alt_klasor / "cekilmeyen_maclar.xlsx",
        "checkpoint": alt_klasor / "ilerleme_kaydi.json",
        "log": alt_klasor / "ilerleme_log.jsonl",
        "mac_listesi_gecici": alt_klasor / "mac_listesi_gecici.json",
    }


def _gun_araligi(basla, son):
    baslangic = datetime.strptime(basla, "%Y-%m-%d").date()
    bitis = datetime.strptime(son, "%Y-%m-%d").date()
    gunler = []
    g = baslangic
    while g <= bitis:
        gunler.append(g.strftime("%Y-%m-%d"))
        g += timedelta(days=1)
    return gunler


# ═════════════════════════════════════════════════════════════
# v15.7 KRİTİK DÜZELTME #2:
#   Gün "tamam" sayılmak için SADECE `tamamlanan_gunler` kümesine
#   güvenilir. `mac_dates` (bir maç bulundu) kanıt DEĞİLDİR.
# ═════════════════════════════════════════════════════════════
def _eksik_gunleri_bul(basla, son, tamamlanan_gunler=None):
    if tamamlanan_gunler is None:
        tamamlanan_gunler = set()
    return [g for g in _gun_araligi(basla, son) if g not in tamamlanan_gunler]


def _tamamlanan_gunleri_guncelle(liste_hatalari, gunler, tamamlanan_gunler):
    """
    Hata ALMAYAN günleri tamamlanmış olarak işaretle.
    Bir gün için 'tarih:YYYY-MM-DD' anahtarı hatalar içinde varsa
    o gün tamamlanmamış sayılır (yeniden denenmeli).
    """
    hata_dict = hata_listesini_dict_yap(liste_hatalari)
    for gun in gunler:
        if f"tarih:{gun}" not in hata_dict:
            tamamlanan_gunler.add(gun)


# ============================================================
# YARDIMCILAR (devam)
# ============================================================

def arsiv_slug_yap(metin):
    metin = metin_temizle(metin)
    cevrim = {
        "ı": "i", "İ": "I", "ğ": "g", "Ğ": "G",
        "ü": "u", "Ü": "U", "ş": "s", "Ş": "S",
        "ö": "o", "Ö": "O", "ç": "c", "Ç": "C",
    }
    metin = metin.translate(str.maketrans(cevrim))
    metin = re.sub(r"[^A-Za-z0-9]+", "-", metin)
    return metin.strip("-")


def oran_temizle(deger):
    deger = metin_temizle(deger)
    if deger.lower() in {"", "-", "–", "—", "none", "nan", "null"}:
        return ""
    deger = deger.replace(" ", "")
    if not re.fullmatch(r"\d+(?:[.,]\d+)?", deger):
        return ""
    return deger.replace(".", ",")


def sayiya_cevir(deger):
    if deger is None or isinstance(deger, bool):
        return None
    try:
        metin = str(deger).strip()
        if not metin or metin.lower() in {"none", "null", "-"}:
            return None
        return int(float(metin.replace(",", ".")))
    except (TypeError, ValueError, OverflowError):
        return None


def bahis_basligi(bahis_turu, secim):
    market = metin_temizle(bahis_turu) or "Bahis türü belirtilmemiş"
    secim = metin_temizle(secim) or "Seçim belirtilmemiş"
    return f"{market} | {secim}"


def oran_kolonu_ekle(oran_kolonlari, kolon):
    if kolon and kolon not in MAC_KOLONLARI and kolon not in oran_kolonlari:
        oran_kolonlari.append(kolon)


def oran_kolonlarini_tamamla(oran_kolonlari, genis_veri):
    for bilgi in genis_veri.values():
        if not isinstance(bilgi, dict):
            continue
        for kolon in bilgi.keys():
            if kolon.startswith("__"):
                continue
            if kolon not in MAC_KOLONLARI and not str(kolon).startswith(ISTATISTIK_ON_EK):
                oran_kolonu_ekle(oran_kolonlari, kolon)


def _bilgiden_oran_istatistik_sil(bilgi, oran_kolonlari):
    for kolon in list(bilgi.keys()):
        if kolon in MAC_KOLONLARI:
            continue
        if str(kolon).startswith("__"):
            continue
        if str(kolon).startswith(ISTATISTIK_ON_EK):
            bilgi.pop(kolon, None)
        elif kolon in oran_kolonlari:
            bilgi.pop(kolon, None)


def _strip_internal_fields(bilgi):
    if not isinstance(bilgi, dict):
        return {}
    return {k: v for k, v in bilgi.items() if not k.startswith("__")}


# ============================================================
# JSON YAZMA
# ============================================================

def _guvenli_json_yaz(hedef, veri):
    hedef = Path(hedef)
    gecici = hedef.with_suffix(".tmp")
    with open(gecici, "w", encoding="utf-8") as f:
        json.dump(veri, f, ensure_ascii=False, separators=(",", ":"))
        f.flush()
        os.fsync(f.fileno())
    os.replace(gecici, hedef)


# ============================================================
# APPEND-ONLY LOG — thread-safe
# ============================================================

class CheckpointLog:
    def __init__(self, log_yolu):
        self.log_yolu = Path(log_yolu)
        self._fh = None
        self._lock = threading.Lock()

    def _ac_kilitli(self):
        if self._fh is None:
            self._fh = open(self.log_yolu, "a", encoding="utf-8")

    def ac(self):
        with self._lock:
            self._ac_kilitli()

    def kapat(self):
        with self._lock:
            if self._fh is not None:
                try:
                    self._fh.close()
                except Exception:
                    pass
                self._fh = None

    def mac_yaz(self, mac_id, bilgi, hata_listesi):
        if not LOG_AKTIF:
            return
        entry = {
            "id": str(mac_id),
            "bilgi": _strip_internal_fields(bilgi),
            "hata": list(hata_listesi) if hata_listesi else None,
        }
        satir = json.dumps(entry, ensure_ascii=False, separators=(",", ":")) + "\n"
        with self._lock:
            self._ac_kilitli()
            try:
                self._fh.write(satir)
                self._fh.flush()
            except Exception as e:
                print(f"[LOG YAZMA HATASI] {mac_id}: {e}")

    def temizle(self):
        with self._lock:
            if self._fh is not None:
                try:
                    self._fh.close()
                except Exception:
                    pass
                self._fh = None
            if self.log_yolu.exists():
                try:
                    self.log_yolu.unlink()
                except Exception as hata:
                    print(f"[UYARI] Log silinemedi: {hata}")


def _logu_uygula(log_yolu, genis_veri, hatalar):
    log_yolu = Path(log_yolu)
    if not log_yolu.exists():
        return 0
    sayi = 0
    bozuk = 0
    with open(log_yolu, "r", encoding="utf-8") as f:
        for satir in f:
            satir = satir.strip()
            if not satir:
                continue
            try:
                entry = json.loads(satir)
            except json.JSONDecodeError:
                bozuk += 1
                continue
            mac_id = entry.get("id")
            if not mac_id:
                continue
            if isinstance(entry.get("bilgi"), dict):
                genis_veri[str(mac_id)] = entry["bilgi"]
            key = f"id:{mac_id}"
            hata = entry.get("hata")
            # v15.7: yapı doğrulaması
            if isinstance(hata, list):
                temiz = [k for k in hata if isinstance(k, dict)]
                if temiz:
                    hatalar[key] = temiz
                else:
                    hatalar.pop(key, None)
            elif hata is None:
                hatalar.pop(key, None)
            else:
                print(f"[UYARI] Log'da beklenmedik hata tipi: "
                      f"{type(hata).__name__} (id={mac_id})")
            sayi += 1
    if bozuk:
        print(f"[UYARI] Log'da {bozuk} bozuk satır atlandı.")
    return sayi


def checkpoint_yukle(dosya, basla, son, log_dosya=None):
    if not dosya.exists():
        return None
    try:
        with open(dosya, "r", encoding="utf-8") as dosya_obj:
            veri = json.load(dosya_obj)
        if (veri.get("baslangic_tarihi") != basla
                or veri.get("bitis_tarihi") != son):
            print("[UYARI] Ara kayıt tarih aralığı farklı; kullanılmayacak.")
            return None
        veri["hatalar"] = hata_listesini_dict_yap(veri.get("hatalar", []))
        veri["tamamlanan_gunler"] = set(veri.get("tamamlanan_gunler", []))
        veri["kalici_hata_idler"] = set(veri.get("kalici_hata_idler", []) or [])
        genis_veri = {str(k): v for k, v in veri.get("genis_veri", {}).items()}
        veri["genis_veri"] = genis_veri
        if log_dosya:
            eklenen = _logu_uygula(log_dosya, genis_veri, veri["hatalar"])
            if eklenen:
                print(f"[DEVAM] Log'dan {eklenen:,} maç güncellemesi uygulandı.")
        print(f"[DEVAM] Önceki ara kayıt yüklendi: "
              f"{veri.get('kayit_zamani', 'zaman bilgisi yok')}")
        if veri["kalici_hata_idler"]:
            print(f"[DEVAM] Önceki oturumdan {len(veri['kalici_hata_idler']):,} "
                  f"kalıcı 401/403 ID yüklendi.")
        return veri
    except Exception as hata:
        print(f"[UYARI] JSON ara kaydı okunamadı: {hata}")
        return None


def checkpoint_yaz(
    yollar, basla, son, genis_veri, hatalar, oran_kolonlari,
    mac_listesi=None, mac_listesi_durum=None, tamamlanan_gunler=None,
    kalici_hata_idler=None
):
    temiz_genis_veri = {}
    for k, v in genis_veri.items():
        temiz_genis_veri[k] = _strip_internal_fields(v) if isinstance(v, dict) else v

    veri = {
        "baslangic_tarihi": basla,
        "bitis_tarihi": son,
        "kayit_zamani": datetime.now().astimezone().isoformat(timespec="seconds"),
        "lig_filtresi": _lig_filtresi_imzasi(),
        "aktif_gruplar": AKTIF_GRUPLAR,
        "oran_kolonlari": oran_kolonlari,
        "mac_listesi": mac_listesi,
        "mac_listesi_durum": mac_listesi_durum,
        "tamamlanan_gunler": sorted(tamamlanan_gunler) if tamamlanan_gunler else [],
        "kalici_hata_idler": sorted(kalici_hata_idler) if kalici_hata_idler else [],
        "genis_veri": temiz_genis_veri,
        "hatalar": hatalar,
    }
    _guvenli_json_yaz(yollar["checkpoint"], veri)


def _checkpoint_mac_listesi_dolu(yollar):
    try:
        with open(yollar["checkpoint"], "r", encoding="utf-8") as f:
            ck = json.load(f)
        ml = ck.get("mac_listesi")
        return isinstance(ml, list) and len(ml) > 0
    except Exception:
        return False


# ============================================================
# HTTP OTURUMU (senkron — maç listesi için)
# ============================================================

class KaliciHata(RuntimeError):
    pass


class ErisimKisitlamaHatasi(RuntimeError):
    def __init__(self, mesaj, url, durum_kodu=429, retry_after_saniye=None):
        super().__init__(mesaj)
        self.url = url
        self.durum_kodu = durum_kodu
        self.retry_after_saniye = retry_after_saniye


_oturum_kilidi = threading.Lock()
_oturum = None


def oturum_al():
    global _oturum
    if _oturum is None:
        with _oturum_kilidi:
            if _oturum is None:
                _oturum = cloudscraper.create_scraper(
                    browser={"browser": "chrome", "platform": "windows",
                             "desktop": True})
    return _oturum


def _retry_after_saniye(cevap):
    try:
        headers = cevap.headers
        ham = headers.get("Retry-After") or headers.get("retry-after") or ""
        ham = str(ham).strip()
    except Exception:
        return None
    if not ham:
        return None
    if ham.isdigit():
        return max(1, min(int(ham), RETRY_AFTER_BILGI_MAKS))
    try:
        hedef = parsedate_to_datetime(ham)
        if hedef.tzinfo is None:
            hedef = hedef.replace(tzinfo=timezone.utc)
        fark = (hedef - datetime.now(timezone.utc)).total_seconds()
        if fark > 0:
            return max(1, min(int(fark), RETRY_AFTER_BILGI_MAKS))
    except (TypeError, ValueError, OverflowError):
        pass
    return None


def istek_yap(url, params=None, livedata=False):
    deneme_5xx = 0
    deneme_ag = 0
    while True:
        try:
            cevap = oturum_al().get(
                url, params=params, timeout=ZAMAN_ASIMI, headers=ISTEK_HEADERS)
            durum = cevap.status_code
            if durum == 429:
                retry_after = _retry_after_saniye(cevap)
                raise ErisimKisitlamaHatasi(
                    f"HTTP 429 — URL: {url}. Retry-After={retry_after or 'yok'}",
                    url=url, durum_kodu=429, retry_after_saniye=retry_after)
            if durum in (401, 403):
                raise ErisimKisitlamaHatasi(
                    f"HTTP {durum} — URL: {url}.",
                    url=url, durum_kodu=durum, retry_after_saniye=None)
            if 500 <= durum <= 599:
                deneme_5xx += 1
                if deneme_5xx >= DENEME_SAYISI_5XX:
                    raise RuntimeError(
                        f"HTTP {durum}; {DENEME_SAYISI_5XX} deneme yetersiz: {url}")
                bekle = min((30 * (2 ** (deneme_5xx - 1))) + random.uniform(0, 5), 300)
                print(f"[HTTP {durum}] {url[:80]} | {sure_formatla(bekle)} bekle")
                time.sleep(bekle)
                continue
            if 400 <= durum <= 499:
                raise KaliciHata(f"HTTP {durum} - {url}")
            cevap.raise_for_status()
            return cevap
        except (ErisimKisitlamaHatasi, KaliciHata):
            raise
        except requests.exceptions.RequestException as hata:
            deneme_ag += 1
            if deneme_ag >= DENEME_SAYISI_AG:
                raise RuntimeError(f"Bağlantı hatası: {url} | {hata}") from hata
            bekle = (15 * (2 ** (deneme_ag - 1))) + random.uniform(0, 3)
            print(f"[AĞ] {url[:60]} | {sure_formatla(bekle)} bekle: {hata}")
            time.sleep(bekle)


# ============================================================
# ASENKRON HTTP — WARP OTO-RESET (kilit yalnızca throttle'da)
# ============================================================

_http_lock = None
_global_cooldown_until = 0.0
_last_request_time = 0.0  # planlanan istek yuvası
_last_actual_request_time = 0.0  # fiilen başlatılan son istek
_consecutive_429_global = 0
_warp_reset_in_progress = False


def _lock_al():
    global _http_lock
    if _http_lock is None:
        _http_lock = asyncio.Lock()
    return _http_lock


class _AsyncCevap:
    __slots__ = ("status_code", "headers", "text", "_json_cache")

    def __init__(self, status_code, headers, text):
        self.status_code = status_code
        try:
            self.headers = {str(k).lower(): v for k, v in headers.items()}
        except Exception:
            self.headers = {}
        self.text = text or ""
        self._json_cache = None

    def json(self):
        if self._json_cache is None:
            self._json_cache = json.loads(self.text)
        return self._json_cache


def _warp_cli_bul():
    for aday in WARP_PATH_CANDIDATES:
        if aday == "warp-cli":
            try:
                r = subprocess.run([aday, "--version"], capture_output=True,
                                    timeout=5, text=True)
                if r.returncode == 0:
                    return aday
            except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
                continue
        elif Path(aday).exists():
            return aday
    return None


def _warp_reset_sync():
    warp = _warp_cli_bul()
    if not warp:
        print("[WARP HATA] warp-cli bulunamadı; WARP reset atlanıyor.")
        return False
    try:
        print("\n[WARP] Bağlantı kesiliyor...")
        r = subprocess.run([warp, "disconnect"], capture_output=True,
                            timeout=15, text=True)
        if r.returncode != 0:
            print(f"[WARP HATA] disconnect (kod {r.returncode}): "
                  f"{r.stderr.strip()}")
            return False
        time.sleep(2)
        print("[WARP] Yeni IP alınıyor (Bağlanılıyor)...")
        r = subprocess.run([warp, "connect"], capture_output=True,
                            timeout=15, text=True)
        if r.returncode != 0:
            print(f"[WARP HATA] connect (kod {r.returncode}): "
                  f"{r.stderr.strip()}")
            return False
        time.sleep(3)
        print("[WARP] Bağlantı yenilendi.")
        return True
    except subprocess.TimeoutExpired:
        print("[WARP HATA] Komut zaman aşımına uğradı.")
        return False
    except Exception as e:
        print(f"[WARP HATA] {e}")
        return False


# ═════════════════════════════════════════════════════════════
# v15.7 KRİTİK DÜZELTME #4:
#   HTTP çağrısı artık kilit DIŞINDA. Kilit yalnızca throttle
#   kararını ve 429/WARP durum değişkenlerini koruyor.
#   Böylece ESZAMANLI_ISTEK = 2 gerçekten işliyor.
# ═════════════════════════════════════════════════════════════
async def istek_yap_async(url, params=None, livedata=False, max_retries=6):
    global _global_cooldown_until, _last_request_time
    global _last_actual_request_time
    global _consecutive_429_global, _warp_reset_in_progress

    attempt = 0
    warp_resets = 0

    while attempt < max_retries:
        # Throttle rezervasyonu kısa kilitle yapılır; bekleme kilit DIŞINDADIR.
        while True:
            bekleme = 0.0
            rezerve_edildi = False
            async with _lock_al():
                simdi = time.time()
                if _warp_reset_in_progress:
                    bekleme = 1.0
                elif simdi < _global_cooldown_until:
                    bekleme = min(_global_cooldown_until - simdi, 3.0)
                    if bekleme > GLOBAL_COOLDOWN_LOG_ESIGI:
                        print(f"[GLOBAL COOLDOWN] {_global_cooldown_until - simdi:.0f}s kaldı, bekleniyor...")
                else:
                    istek_zamani = max(simdi, _last_request_time + MIN_ISTEK_ARASI)
                    istek_zamani += random.uniform(0.3, max(0.3, ISTEK_JITTER_ASYNC))
                    _last_request_time = istek_zamani
                    bekleme = max(0.0, istek_zamani - simdi)
                    rezerve_edildi = True
            if bekleme > 0:
                await asyncio.sleep(bekleme)
            if rezerve_edildi:
                ikinci_bekleme = 0.0
                async with _lock_al():
                    simdi = time.time()
                    if _warp_reset_in_progress:
                        rezerve_edildi = False
                        ikinci_bekleme = 1.0
                    elif simdi < _global_cooldown_until:
                        rezerve_edildi = False
                        ikinci_bekleme = min(_global_cooldown_until - simdi, 3.0)
                    else:
                        # Event loop uzun süre bloklanırsa birden çok rezervasyon
                        # aynı anda uyanabilir. Fiilî başlangıç saati bu çakışmayı önler.
                        gecen_gercek = simdi - _last_actual_request_time
                        if gecen_gercek < MIN_ISTEK_ARASI:
                            ikinci_bekleme = (MIN_ISTEK_ARASI - gecen_gercek
                                              + random.uniform(0.3, max(0.3, ISTEK_JITTER_ASYNC)))
                            rezerve_edildi = False
                        else:
                            _last_actual_request_time = simdi
                if rezerve_edildi:
                    break
                if ikinci_bekleme > 0:
                    await asyncio.sleep(ikinci_bekleme)

        # 2. HTTP — kilit dışında (gerçek eşzamanlılık)
        headers = {
            "Referer": "https://arsiv.mackolik.com/",
            "Accept-Language": "tr-TR,tr;q=0.9,en;q=0.8",
        }
        try:
            async with AsyncSession(impersonate=CURL_IMPERSONATE) as session:
                r = await session.get(
                    url, params=params, headers=headers,
                    timeout=60.0, allow_redirects=True,
                )
                wrapped = _AsyncCevap(status_code=r.status_code,
                                       headers=r.headers, text=r.text)
        except Exception as e:
            if _warp_reset_in_progress:
                print(f"[AĞ] {url[:60]} | WARP reset devam ediyor, bekleniyor...")
                await asyncio.sleep(5)
                continue
            print(f"[AĞ] {url[:60]} | {type(e).__name__}: {e}")
            attempt += 1
            if attempt < max_retries:
                await asyncio.sleep((2 ** (attempt - 1)) * 2)
            continue

        durum = wrapped.status_code

        # 3. 429 — durum değişkenleri kilit altında
        if durum == 429:
            should_reset = False
            async with _lock_al():
                if not _warp_reset_in_progress:
                    _consecutive_429_global += 1
                    print(f"[429] {url[:60]} | Birikimli 429: "
                          f"{_consecutive_429_global}/{ARDISIK_429_ESIK}")
                    if _consecutive_429_global >= ARDISIK_429_ESIK:
                        should_reset = True
                        _warp_reset_in_progress = True
                        _consecutive_429_global = 0
                        _global_cooldown_until = max(
                            _global_cooldown_until,
                            time.time() + WARP_RESET_SONRASI_BEKLEME + 60)

            if should_reset:
                warp_resets += 1
                if warp_resets > WARP_RESET_MAKS:
                    print(f"[WARP] Reset limiti ({WARP_RESET_MAKS}) aşıldı; "
                          f"bu istek için tekrar denenmeyecek.")
                    async with _lock_al():
                        _warp_reset_in_progress = False
                        _consecutive_429_global = 0
                    raise ErisimKisitlamaHatasi(
                        f"WARP reset limiti ({WARP_RESET_MAKS}) aşıldı: {url}",
                        url=url, durum_kodu=429, retry_after_saniye=None)

                print(f"\n{'!' * 60}")
                print(f"{ARDISIK_429_ESIK} KEZ 429 ALINDI! "
                      f"OTOMATİK WARP RESET ({warp_resets}/{WARP_RESET_MAKS})...")
                print("!" * 60)
                reset_ok = False
                try:
                    reset_ok = await asyncio.to_thread(_warp_reset_sync)
                except Exception as e:
                    print(f"[WARP HATA] {e}")

                if reset_ok:
                    async with _lock_al():
                        _global_cooldown_until = (time.time()
                                                   + WARP_RESET_SONRASI_BEKLEME)
                        _last_request_time = time.time()
                    print(f"\n[SOĞUMA] {WARP_RESET_SONRASI_BEKLEME} saniye bekleniyor...")
                    bitis = time.time() + WARP_RESET_SONRASI_BEKLEME
                    while True:
                        kalan = bitis - time.time()
                        if kalan <= 0:
                            break
                        print(f"[SOĞUMA] Kalan: {int(kalan)} saniye...")
                        await asyncio.sleep(min(kalan, 30))
                    print("[SOĞUMA] Bitti. Kaldığı yerden devam ediliyor...")
                else:
                    async with _lock_al():
                        _consecutive_429_global = ARDISIK_429_ESIK
                        _global_cooldown_until = max(
                            _global_cooldown_until, time.time() + 120)
                    print("[WARP] Reset başarısız; normal 429 backoff (120s).")

                async with _lock_al():
                    _warp_reset_in_progress = False
                continue

            if _warp_reset_in_progress:
                await asyncio.sleep(5)
                continue

            retry_after = _retry_after_saniye(wrapped)
            tavan = LIVEDATA_RETRY_AFTER_MAX if livedata else RETRY_AFTER_BEKLEME_MAKS
            if retry_after:
                if retry_after <= tavan:
                    wait = retry_after
                else:
                    wait = tavan
                    print(f"[429] Sunucu {retry_after}s istedi; politika "
                          f"tavanı {tavan}s uygulanıyor.")
            else:
                wait = min((2 ** attempt) * 5, 300)
            async with _lock_al():
                _global_cooldown_until = max(_global_cooldown_until,
                                              time.time() + wait)
            print(f"[429 GLOBAL] {url[:60]} | tüm istekler {wait}s duracak "
                  f"({attempt + 1}/{max_retries})")
            attempt += 1
            continue

        # 4. Diğer HTTP hataları
        if durum in (500, 502, 503, 504):
            wait = min((2 ** attempt) * 3, 120)
            print(f"[{durum}] {url[:60]} | {wait}s bekle")
            await asyncio.sleep(wait)
            attempt += 1
            continue

        if durum in (401, 403):
            raise ErisimKisitlamaHatasi(
                f"HTTP {durum} — URL: {url}.",
                url=url, durum_kodu=durum, retry_after_saniye=None)

        if 400 <= durum <= 499:
            raise KaliciHata(f"HTTP {durum} - {url}")

        # 5. Başarılı
        async with _lock_al():
            _consecutive_429_global = 0
        return wrapped

    raise RuntimeError(f"Maks deneme aşıldı: {url[:80]}")


# ============================================================
# MATCHDATA API
# ============================================================

async def matchdata_cek_async(mac_id):
    url = f"https://arsiv.mackolik.com/Match/MatchData.aspx?t=dtl&id={mac_id}&s=0"
    cevap = await istek_yap_async(url, livedata=True)
    metin = cevap.text.strip()
    if not metin or metin == "" or metin == "null":
        raise ValueError(f"MatchData boş yanıt: {url}")
    try:
        veri = json.loads(metin)
    except json.JSONDecodeError as hata:
        raise ValueError(f"MatchData JSON parse hatası: {hata}")
    if not isinstance(veri, dict):
        raise ValueError(f"MatchData dict değil: {type(veri)}")
    return veri


def matchdata_kadro_parse(veri):
    sonuc = {"Ev Kadro": "", "Deplasman Kadro": "", "Ev Yedek": "", "Deplasman Yedek": ""}

    def _oyuncu_str(kayit):
        if not isinstance(kayit, list) or len(kayit) < 2:
            return ""
        ad = metin_temizle(kayit[1])
        forma = ""
        if len(kayit) > 2:
            try:
                if kayit[2] and int(kayit[2]) > 0:
                    forma = str(int(kayit[2]))
            except (TypeError, ValueError):
                forma = ""
        return f"{forma} {ad}" if forma else ad

    def _liste_str(kayitlar, baslangic=0, bitis=None):
        if not isinstance(kayitlar, list):
            return ""
        parcalar = []
        for k in kayitlar[baslangic:bitis]:
            s = _oyuncu_str(k)
            if s:
                parcalar.append(s)
        return "; ".join(parcalar)

    ev = veri.get("h") or []
    dep = veri.get("a") or []
    if ev:
        sonuc["Ev Kadro"] = _liste_str(ev, 0, 11)
        sonuc["Ev Yedek"] = _liste_str(ev, 11)
    if dep:
        sonuc["Deplasman Kadro"] = _liste_str(dep, 0, 11)
        sonuc["Deplasman Yedek"] = _liste_str(dep, 11)
    return sonuc


def matchdata_events_parse(events):
    sonuc = {"Goller": "", "Sarı Kartlar": "", "Kırmızı Kartlar": "",
             "Değişiklikler": "", "Kaçan Penaltılar": ""}
    if not isinstance(events, list):
        return sonuc
    goller, sari, kirmizi, degisiklik, kacan_pen = [], [], [], [], []
    for ev in events:
        if not isinstance(ev, list) or len(ev) < 5:
            continue
        try:
            takim = int(ev[0]); dakika = ev[1]
            oyuncu_adi = metin_temizle(ev[3]) if len(ev) > 3 else ""
            tip = int(ev[4])
        except (TypeError, ValueError):
            continue
        detay = ev[5] if len(ev) > 5 and isinstance(ev[5], dict) else {}
        detay_tipi = sayiya_cevir(detay.get("d"))
        taraf = "Ev" if takim == 1 else "Dep"
        dk_str = f"{dakika}'" if dakika not in (None, "") else "?"
        if tip == 1:
            if not oyuncu_adi and detay.get("p"):
                oyuncu_adi = metin_temizle(detay.get("p"))
            if not oyuncu_adi:
                oyuncu_adi = "Bilinmeyen"
            not_parca = ""
            if detay_tipi == 3:
                not_parca = " (k.k.)"
            elif detay_tipi == 2:
                not_parca = " (P)"
            asist = (f" (asist: {metin_temizle(detay['astName'])})"
                     if detay.get("astName") else "")
            goller.append(f"{dk_str} {oyuncu_adi}{not_parca}{asist} [{taraf}]")
        elif tip == 2:
            sari.append(f"{dk_str} {oyuncu_adi or '?'} [{taraf}]")
        elif tip == 3:
            if detay_tipi == 1:
                kirmizi.append(f"{dk_str} {oyuncu_adi or '?'} (çift sarı) [{taraf}]")
            else:
                kirmizi.append(f"{dk_str} {oyuncu_adi or '?'} [{taraf}]")
        elif tip == 4:
            giren = metin_temizle(detay.get("sub", "")) if detay.get("sub") else ""
            if oyuncu_adi and giren:
                degisiklik.append(f"{dk_str} {oyuncu_adi} → {giren} [{taraf}]")
            elif oyuncu_adi or giren:
                degisiklik.append(f"{dk_str} {oyuncu_adi or '?'} ↔ {giren or '?'} [{taraf}]")
        elif tip == 7:
            kacan_pen.append(f"{dk_str} {oyuncu_adi or '?'} [{taraf}]")
    sonuc["Goller"] = "; ".join(goller)
    sonuc["Sarı Kartlar"] = "; ".join(sari)
    sonuc["Kırmızı Kartlar"] = "; ".join(kirmizi)
    sonuc["Değişiklikler"] = "; ".join(degisiklik)
    sonuc["Kaçan Penaltılar"] = "; ".join(kacan_pen)
    return sonuc


async def _matchdata_uygula_async(bilgi, mac_id, hatalar):
    """YERİNDE mutasyon; (bilgi, hatalar, basarili) döndürür."""
    try:
        veri = await matchdata_cek_async(mac_id)
        kadro = matchdata_kadro_parse(veri)
        event = matchdata_events_parse(veri.get("e", []))
        for k, v in kadro.items():
            bilgi[k] = v
        for k, v in event.items():
            bilgi[k] = v
        hatalar = hata_kayitlarini_temizle(
            hatalar, mac_id, baslangiclar=("MATCHDATA:",))
        return bilgi, hatalar, True
    except ErisimKisitlamaHatasi as hata:
        hata_ekle(hatalar, {
            "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
            "Maç": bilgi.get("Maç", ""),
            "Hata": f"MATCHDATA: HTTP {hata.durum_kodu}; URL={hata.url}",
        })
        return bilgi, hatalar, False
    except Exception as hata:
        hata_ekle(hatalar, {
            "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
            "Maç": bilgi.get("Maç", ""),
            "Hata": f"MATCHDATA: {hata}",
        })
        return bilgi, hatalar, False


# ============================================================
# ANA HTML'DEN KADRO VE OLAYLAR
# ============================================================

def ana_html_den_kadro_ve_olaylar(soup):
    sonuc = {
        "Ev Kadro": "", "Ev Yedek": "",
        "Deplasman Kadro": "", "Deplasman Yedek": "",
        "Goller": "", "Sarı Kartlar": "", "Kırmızı Kartlar": "",
        "Değişiklikler": "", "Kaçan Penaltılar": ""
    }

    def _parse_team(div_id):
        div = soup.select_one(f"#{div_id}")
        if not div:
            return [], []
        starters, subs = [], []
        is_subs = False
        for tr in div.select("tr"):
            classes = tr.get("class", [])
            if "team-value" in classes or "team-value-2" in classes:
                is_subs = True
                continue
            number_el = tr.select_one("td.bld")
            name_el = tr.select_one("td a")
            if name_el:
                name = name_el.text.strip()
                number = number_el.text.strip() if number_el else ""
                player = f"{number} {name}" if number and not number.startswith('0') else name
                (subs if is_subs else starters).append(player)
        return starters, subs

    ev_ilk11, ev_yedek = _parse_team("dvHomeSquad")
    dep_ilk11, dep_yedek = _parse_team("dvAwaySquad")

    sonuc["Ev Kadro"] = "; ".join(ev_ilk11)
    sonuc["Ev Yedek"] = "; ".join(ev_yedek)
    sonuc["Deplasman Kadro"] = "; ".join(dep_ilk11)
    sonuc["Deplasman Yedek"] = "; ".join(dep_yedek)

    div_events = soup.select_one("#dvEvents")
    if div_events:
        goller, sari, kirmizi, degisiklik, kacan_pen = [], [], [], [], []
        for tr in div_events.select("tr"):
            time_el = tr.select_one(".list-table-time")
            if not time_el:
                continue
            dk = time_el.text.strip()
            tds = tr.find_all("td")
            if len(tds) >= 7:
                home_text = tds[0].get_text(" ", strip=True)
                home_icon = tds[1].find("img")
                score_text = tds[2].get_text(" ", strip=True)
                away_icon = tds[5].find("img")
                away_text = tds[6].get_text(" ", strip=True)

                if home_icon:
                    title = home_icon.get("title")
                    if title in ["Gol", "Penaltıdan Gol"]:
                        goller.append(f"{dk} {home_text} {score_text} [Ev]")
                    elif title == "Sarı Kart":
                        sari.append(f"{dk} {home_text} [Ev]")
                    elif title in ["Kırmızı Kart", "Çift Sarıdan Kırmızı Kart"]:
                        kirmizi.append(f"{dk} {home_text} [Ev]")
                    elif title == "Giren Oyuncu":
                        degisiklik.append(f"{dk} {home_text} [Ev]")
                    elif title == "Kaçan Penaltı":
                        kacan_pen.append(f"{dk} {home_text} [Ev]")

                if away_icon:
                    title = away_icon.get("title")
                    if title in ["Gol", "Penaltıdan Gol"]:
                        goller.append(f"{dk} {away_text} {score_text} [Dep]")
                    elif title == "Sarı Kart":
                        sari.append(f"{dk} {away_text} [Dep]")
                    elif title in ["Kırmızı Kart", "Çift Sarıdan Kırmızı Kart"]:
                        kirmizi.append(f"{dk} {away_text} [Dep]")
                    elif title == "Giren Oyuncu":
                        degisiklik.append(f"{dk} {away_text} [Dep]")
                    elif title == "Kaçan Penaltı":
                        kacan_pen.append(f"{dk} {away_text} [Dep]")

        sonuc["Goller"] = "; ".join(goller)
        sonuc["Sarı Kartlar"] = "; ".join(sari)
        sonuc["Kırmızı Kartlar"] = "; ".join(kirmizi)
        sonuc["Değişiklikler"] = "; ".join(degisiklik)
        sonuc["Kaçan Penaltılar"] = "; ".join(kacan_pen)

    return sonuc


# ============================================================
# ARŞİV SAYFASI PARSE
# ============================================================

def mac_arsiv_url(row):
    match_id = metin_temizle(row.get("Maç ID", ""))
    ev = metin_temizle(row.get("Ev Sahibi", ""))
    dep = metin_temizle(row.get("Deplasman", ""))
    if not match_id or not match_id.isdigit():
        raise ValueError(f"Sayısal arşiv maç ID'si yok: {match_id!r}")
    if not ev or not dep:
        raise ValueError(f"Takım adları eksik: ID={match_id}")
    slug = f"{arsiv_slug_yap(ev)}-{arsiv_slug_yap(dep)}"
    return f"https://arsiv.mackolik.com/Mac/{match_id}/{slug}"


def _skor_duzenle(deger):
    metin = metin_temizle(deger)
    match = re.search(r"(?<!\d)(\d{1,2})\s*[-–—:]\s*(\d{1,2})(?!\d)", metin)
    if not match:
        return ""
    return f"{int(match.group(1))}-{int(match.group(2))}"


def ilk_yari_skoru_scriptlerde_ara(soup):
    desenler = [
        ("htScore alanı", re.compile(
            r"[\"']?(?P<alan>htScore|halfTimeScore|halftimeScore|htscore)"
            r"[\"']?\s*[:=]\s*[\"'](?P<skor>\d{1,2}\s*[-–:：]\s*\d{1,2})[\"']",
            flags=re.IGNORECASE)),
        ("#dvHTScoreText DOM", re.compile(
            r"(?:getElementById\(\s*[\"']dvHTScoreText[\"']\s*\)"
            r"|querySelector\(\s*[\"']#dvHTScoreText[\"']\s*\))"
            r"[\s\S]{0,160}?(?:innerHTML|textContent|innerText)\s*=\s*"
            r"[\"'](?P<skor>\d{1,2}\s*[-–:：]\s*\d{1,2})[\"']",
            flags=re.IGNORECASE)),
        ("#dvHTScoreText jQuery", re.compile(
            r"\$\(\s*[\"']#dvHTScoreText[\"']\s*\)\s*\.\s*(?:text|html)\("
            r"\s*[\"'](?P<skor>\d{1,2}\s*[-–:：]\s*\d{1,2})[\"']\s*\)",
            flags=re.IGNORECASE)),
    ]
    for script in soup.find_all("script"):
        metin = script.string or script.get_text(" ", strip=True) or ""
        for kaynak, desen in desenler:
            eslesen = desen.search(metin)
            if not eslesen:
                continue
            ham = eslesen.groupdict().get("skor")
            skor = _skor_duzenle(ham)
            if skor:
                alan = eslesen.groupdict().get("alan")
                kaynak_ad = (f"Arşiv JavaScript alanı: {alan}"
                             if alan else f"Arşiv JavaScript: {kaynak}")
                return skor, kaynak_ad
    return "", ""


def arsiv_mac_bilgisi_cikar(soup, row):
    info = soup.select_one(".match-info-wrapper")
    if info is None:
        raise ValueError("Arşiv sayfasında .match-info-wrapper bulunamadı.")
    sonuc = {}
    lig = info.select_one(".match-info-wrapper-season a")
    sonuc["Lig/Sezon"] = metin_temizle(lig.get_text(" ", strip=True)) if lig else ""
    tarih_el = info.select_one(".match-info-date")
    tarih_metin = metin_temizle(tarih_el.get_text(" ", strip=True)) if tarih_el else ""
    tarih_es = re.search(r"(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2}:\d{2})", tarih_metin)
    if tarih_es:
        try:
            dt = datetime.strptime(
                f"{tarih_es.group(1)} {tarih_es.group(2)}", "%d.%m.%Y %H:%M")
            sonuc["Tarih"] = dt.strftime("%d-%m-%Y")
            sonuc["Saat"] = dt.strftime("%H:%M")
        except ValueError:
            sonuc["Tarih"] = metin_temizle(row.get("Tarih", ""))
            sonuc["Saat"] = metin_temizle(row.get("Saat", ""))
    else:
        sonuc["Tarih"] = metin_temizle(row.get("Tarih", ""))
        sonuc["Saat"] = metin_temizle(row.get("Saat", ""))
    ev_el = info.select_one("a.left-block-team-name")
    dep_el = info.select_one("a.r-left-block-team-name")
    ev = (metin_temizle(ev_el.get_text(" ", strip=True))
          if ev_el else metin_temizle(row.get("Ev Sahibi", "")))
    dep = (metin_temizle(dep_el.get_text(" ", strip=True))
           if dep_el else metin_temizle(row.get("Deplasman", "")))
    sonuc["Ev Sahibi"] = ev
    sonuc["Deplasman"] = dep
    sonuc["Maç"] = f"{ev} vs {dep}" if ev and dep else metin_temizle(row.get("Maç", ""))
    sonuc["Maç ID"] = str(row.get("Maç ID", ""))
    sonuc["İddaa Kodu"] = str(row.get("İddaa Kodu", ""))
    durum_el = info.select_one("#dvStatusText")
    sonuc["Maç Durumu"] = metin_temizle(durum_el.get_text(" ", strip=True)) if durum_el else ""
    skor_el = info.select_one("#dvScoreText")
    skor = _skor_duzenle(skor_el.get_text(" ", strip=True)) if skor_el else ""
    sonuc["Skor"] = skor or metin_temizle(row.get("Skor", ""))
    iy_el = info.select_one("#dvHTScoreText")
    iy_skor = _skor_duzenle(iy_el.get_text(" ", strip=True)) if iy_el else ""
    iy_kaynak = "Arşiv HTML #dvHTScoreText" if iy_skor else ""
    if not iy_skor:
        iy_skor, js_kaynak = ilk_yari_skoru_scriptlerde_ara(soup)
        if iy_skor:
            iy_kaynak = js_kaynak
    if not iy_skor:
        iy_skor = _skor_duzenle(row.get("İY Skor", ""))
        if iy_skor:
            iy_kaynak = "Maç listesi ilk yarı skoru"
    sonuc["İY Skor"] = iy_skor
    sonuc["İY Skor Kaynağı"] = iy_kaynak
    seyirci_el = info.select_one(".left-block .hf-match-score")
    seyirci_metin = metin_temizle(seyirci_el.get_text(" ", strip=True)) if seyirci_el else ""
    seyirci_es = re.search(r"Seyirci\s*:\s*(.+)", seyirci_metin, flags=re.IGNORECASE)
    sonuc["Seyirci"] = metin_temizle(seyirci_es.group(1)) if seyirci_es else ""
    stadyum = info.select_one('.match-info-wrapper-bottom a[itemprop="location"]')
    sonuc["Stadyum"] = metin_temizle(stadyum.get_text(" ", strip=True)) if stadyum else ""
    hakem = info.select_one('.match-info-wrapper-bottom a[href*="/Hakem/"]')
    sonuc["Hakem"] = metin_temizle(hakem.get_text(" ", strip=True)) if hakem else ""
    ev_hoca = info.select_one(".left-block .coach a")
    dep_hoca = info.select_one(".right-block .r-coach a")
    sonuc["Ev Sahibi Teknik Direktör"] = (
        metin_temizle(ev_hoca.get_text(" ", strip=True)) if ev_hoca else "")
    sonuc["Deplasman Teknik Direktör"] = (
        metin_temizle(dep_hoca.get_text(" ", strip=True)) if dep_hoca else "")
    alt = info.select_one(".match-info-wrapper-bottom .match-info-detail-text")
    if alt:
        alt_kopya = BeautifulSoup(str(alt), "html.parser")
        for node in alt_kopya.select("a, img, script, style"):
            node.decompose()
        tv = metin_temizle(alt_kopya.get_text(" ", strip=True))
        tv = re.sub(r"^Tv Yayını\s*", "", tv, flags=re.IGNORECASE)
        sonuc["TV Yayını"] = tv.strip(" ,;|")
    else:
        sonuc["TV Yayını"] = ""
    sonuc["Maç Slug"] = ""

    kadro_ve_olaylar = ana_html_den_kadro_ve_olaylar(soup)
    sonuc.update(kadro_ve_olaylar)
    return sonuc


def arsiv_sayfa_oranlari(soup):
    kapsayici = soup.select_one("#compare-left-coll")
    if kapsayici is None:
        return []
    rows = []
    for md in kapsayici.select("div.md"):
        title_div = (md.find("div", class_="detail-title", recursive=False)
                     or md.find("div", class_="detail-title"))
        if title_div is None:
            continue
        secim_divler = md.find_all("div", class_="sgoutcome-name")
        if not secim_divler:
            continue
        parcalar = []
        for child in title_div.children:
            if getattr(child, "name", None) == "span":
                break
            if isinstance(child, str):
                parcalar.append(child)
        market = metin_temizle(" ".join(parcalar)) or "Bahis türü belirtilmemiş"
        for secim_div in secim_divler:
            secim = metin_temizle(secim_div.get_text(" ", strip=True))
            if not secim:
                continue
            outcome = secim_div.parent
            value_div = outcome.find("div", class_="sgoutcome-value") if outcome else None
            if value_div is None:
                continue
            oran = oran_temizle(value_div.get_text(" ", strip=True))
            if oran:
                rows.append({"Bahis Türü": market, "Seçim": secim, "Oran": oran})
    if not rows:
        for tablo in kapsayici.select("table.iddaa-ms-h"):
            _eski_tablo_ayristir(tablo, rows)
    benzersiz, gorulen = [], set()
    for item in rows:
        key = (item["Bahis Türü"], item["Seçim"], item["Oran"])
        if key not in gorulen:
            gorulen.add(key)
            benzersiz.append(item)
    return benzersiz


def _eski_tablo_ayristir(tablo, rows):
    """tbody altındaki satırlar da taranır."""
    baslik_satirlari = tablo.find_all("tr", class_="iddaa-tab-title")
    if not baslik_satirlari:
        return
    baslik_tdler = baslik_satirlari[0].find_all("td", recursive=False)
    kolon_market = {}
    idx = 0
    for td in baslik_tdler:
        ad = metin_temizle(td.get_text(" ", strip=True))
        try:
            cs = int(td.get("colspan", 1))
        except (TypeError, ValueError):
            cs = 1
        for k in range(cs):
            kolon_market[idx + k] = ad
        idx += cs
    tum_tr = tablo.find_all("tr")
    bekleyen = None
    for tr in tum_tr:
        sinif = tr.get("class") or []
        if "iddaa-tab-title" in sinif:
            continue
        if "iddaa-tab-alt" in sinif:
            bekleyen = tr
            continue
        if bekleyen is None:
            continue
        secim_tdler = bekleyen.find_all("td", recursive=False)
        deger_tdler = tr.find_all("td", recursive=False)
        for k, secim_td in enumerate(secim_tdler):
            secim = metin_temizle(secim_td.get_text(" ", strip=True))
            if not secim or k >= len(deger_tdler):
                continue
            a = deger_tdler[k].find("a")
            if a is None:
                continue
            oran = oran_temizle(a.get_text(" ", strip=True))
            if not oran:
                continue
            market = kolon_market.get(k, "")
            if not market:
                continue
            rows.append({"Bahis Türü": market, "Seçim": secim, "Oran": oran})
        bekleyen = None


# ============================================================
# İSTATİSTİK
# ============================================================

ISTATISTIK_ON_TARAF_EK = ISTATISTIK_ON_EK


def _satirlari_parse_et(kapsayici):
    satirlar = kapsayici.select(".match-statistics-rows, .match-statistics-rows-2")
    if not satirlar:
        return {}
    bulunan = {}
    for satir in satirlar:
        isim_el = satir.select_one(".statistics-title-text")
        ev_el = satir.select_one(".team-1-statistics-text")
        dep_el = satir.select_one(".team-2-statistics-text")
        if isim_el is None:
            continue
        isim = metin_temizle(isim_el.get_text(" ", strip=True))
        isim = ISTATISTIK_NORMALIZE.get(isim, isim)
        ev = metin_temizle(ev_el.get_text(" ", strip=True)) if ev_el else ""
        dep = metin_temizle(dep_el.get_text(" ", strip=True)) if dep_el else ""
        if isim and (ev or dep):
            bulunan[f"{ISTATISTIK_ON_TARAF_EK}{isim} | Ev Sahibi"] = ev
            bulunan[f"{ISTATISTIK_ON_TARAF_EK}{isim} | Deplasman"] = dep
    return bulunan


_AJAX_HATA_SINYALLERI = (
    "error", "forbidden", "unauthorized", "exception",
    "hata", "yetki", "erişim", "erişim engellendi",
)


async def _istatistik_ajax_cek_async(mac_id):
    son_hata = None
    for cmd in ("optaStats", "rbStats"):
        try:
            url = (f"https://arsiv.mackolik.com/AjaxHandlers/"
                   f"MatchHandler.aspx?command={cmd}&id={mac_id}")
            cevap = await istek_yap_async(url, livedata=True)
            durum = cevap.status_code
            metin = (cevap.text or "").strip()

            # 204 başarılı ama içeriksiz yanıttır; hata olarak raporlanmaz.
            if durum == 204:
                continue

            # Koruyucu kontrol: HTTP katmanı 206 gibi diğer 2xx kodlarını da döndürebilir.
            if durum != 200:
                son_hata = f"HTTP {durum}"
                print(f"[ISTATISTIK AJAX {cmd}] HTTP {durum} - {mac_id}")
                continue

            if metin.lower().startswith(("<!doctype", "<html")):
                son_hata = f"HTML yanıtı: {metin[:80]!r}"
                continue

            if not metin or metin in ("null", "[]", "{}"):
                continue

            if len(metin) < 100:
                alt = metin.lower()
                if any(k in alt for k in _AJAX_HATA_SINYALLERI):
                    son_hata = f"kısa hata yanıtı: {metin[:80]!r}"
                continue

            soup2 = BeautifulSoup(metin, "html.parser")
            alt = (soup2.select_one("#dvOPTAStats")
                   or soup2.select_one("#dvRBStats") or soup2)
            bulunan = _satirlari_parse_et(alt)
            if bulunan:
                return bulunan, cmd, None
        except ErisimKisitlamaHatasi as hata:
            son_hata = f"HTTP {hata.durum_kodu}"
            print(f"[ISTATISTIK AJAX {cmd}] Erişim hatası "
                  f"{hata.durum_kodu} - {mac_id}")
        except KaliciHata as hata:
            son_hata = f"Kalıcı: {hata}"
        except Exception as hata:
            son_hata = f"{type(hata).__name__}: {hata}"
            print(f"[ISTATISTIK AJAX {cmd} HATASI] {mac_id}: {hata}")
    return {}, None, son_hata


async def istatistikleri_cikar_async(soup, mac_id=None):
    kapsayici = soup.select_one("#dvOPTAStats") or soup.select_one("#dvRBStats")
    if kapsayici is not None:
        bulunan = _satirlari_parse_et(kapsayici)
        if bulunan:
            return bulunan, "bulundu"
    if mac_id:
        bulunan, cmd, hata = await _istatistik_ajax_cek_async(mac_id)
        if bulunan:
            return bulunan, f"bulundu (AJAX {cmd})"
        if hata:
            return {}, f"hata: {hata}"
    return {}, "kapsayici_yok"


def hata_sayfasi_mi(soup):
    title = soup.find("title")
    if title and "hata" in title.get_text(" ", strip=True).lower():
        return True
    return soup.find("h1", class_="error-number") is not None


async def get_bahis_oranlari_async(row):
    url = mac_arsiv_url(row)
    cevap = await istek_yap_async(url)
    soup = BeautifulSoup(cevap.text, "html.parser")
    if hata_sayfasi_mi(soup):
        raise KaliciHata(f"Arşiv hata sayfası döndürdü: {url}")
    if soup.select_one(".match-info-wrapper") is None:
        raise KaliciHata(f"Arşiv maç içeriği bulunamadı: {url}")

    mac_bilgisi = arsiv_mac_bilgisi_cikar(soup, row)
    oranlar = arsiv_sayfa_oranlari(soup)
    mac_id_ver = metin_temizle(row.get("Maç ID", ""))
    istatistikler, ist_durum = await istatistikleri_cikar_async(soup, mac_id=mac_id_ver)

    return {
        "mac_bilgisi": mac_bilgisi,
        "oranlar": oranlar,
        "istatistikler": istatistikler,
        "istatistik_durum": ist_durum,
        "kaynak_url": url,
    }


# ============================================================
# EXCEL
# ============================================================

def _excel_yaz_guvenli(df_maclar, dosya, max_deneme=3):
    hedef = Path(dosya)
    gecici = hedef.with_name(f"{hedef.stem}.tmp{hedef.suffix}")
    for deneme in range(max_deneme):
        try:
            with pd.ExcelWriter(gecici, engine="openpyxl") as writer:
                df_maclar.to_excel(writer, index=False, sheet_name="Maclar")
                ws = writer.sheets["Maclar"]
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for sutun_no in range(1, ws.max_column + 1):
                    hucre = ws.cell(row=1, column=sutun_no)
                    baslik = "" if hucre.value is None else str(hucre.value)
                    ws.column_dimensions[hucre.column_letter].width = min(
                        max(len(baslik) + 2, 12), 45)

            if not gecici.exists() or gecici.stat().st_size < 200:
                boyut = gecici.stat().st_size if gecici.exists() else 0
                raise RuntimeError(f"Geçici Excel geçersiz ({boyut} bayt)")
            try:
                from openpyxl import load_workbook
                _wb = load_workbook(gecici, read_only=True)
                _wb.close()
            except Exception as verify_err:
                raise RuntimeError(f"Geçici Excel doğrulanamadı: {verify_err}")

            for replace_deneme in range(max_deneme):
                try:
                    os.replace(gecici, hedef)
                    return True
                except PermissionError:
                    if replace_deneme < max_deneme - 1:
                        print(f"[EXCEL] Dosya kilitli, "
                              f"{replace_deneme + 1}. deneme 2s sonra...")
                        time.sleep(2)
                    else:
                        raise
            return True
        except PermissionError as e:
            if deneme < max_deneme - 1:
                print(f"[EXCEL] Yazma hatası, {deneme + 1}. deneme 2s sonra: {e}")
                time.sleep(2)
            else:
                print(f"[EXCEL] Yazma başarısız (dosya kilitli): {e}")
                return False
        except Exception as e:
            print(f"[EXCEL HATASI] {e}")
            if gecici.exists():
                try:
                    gecici.unlink()
                except Exception:
                    pass
            return False
    return False


def hata_raporu_yaz(hatalar, dosya):
    hedef = Path(dosya)
    gecici = hedef.with_name(f"{hedef.stem}.tmp{hedef.suffix}")
    hatalar_liste = hatalari_listeye_cevir(hatalar)
    try:
        pd.DataFrame(hatalar_liste, columns=HATA_KOLONLARI).to_excel(
            gecici, index=False, engine="openpyxl")

        try:
            from openpyxl import load_workbook
            wb = load_workbook(gecici, read_only=True)
            wb.close()
        except Exception as verify_err:
            print(f"[HATA RAPORU] Geçici dosya doğrulanamadı: {verify_err}")
            if gecici.exists():
                try:
                    gecici.unlink()
                except Exception:
                    pass
            return False

        for _ in range(3):
            try:
                os.replace(gecici, hedef)
                return True
            except PermissionError:
                time.sleep(2)
        try:
            os.replace(gecici, hedef)
            return True
        except Exception as e:
            print(f"[HATA RAPORU] Değiştirilemedi: {e}")
            return False
    except Exception as e:
        print(f"[HATA RAPORU] {e}")
        if gecici.exists():
            try:
                gecici.unlink()
            except Exception:
                pass
        return False


def excel_ara_kaydet(genis_veri, oran_kolonlari, excel_dosya, hata_dosya, hatalar):
    oranlar = [k for k in oran_kolonlari
               if k not in MAC_KOLONLARI and not str(k).startswith(ISTATISTIK_ON_EK)]
    ist_kolonlari = []
    for bilgi in genis_veri.values():
        if not isinstance(bilgi, dict):
            continue
        for kolon in bilgi.keys():
            if str(kolon).startswith(ISTATISTIK_ON_EK) and kolon not in ist_kolonlari:
                ist_kolonlari.append(kolon)
    tum_kolonlar = MAC_KOLONLARI + oranlar + ist_kolonlari
    df_genis = pd.DataFrame(list(genis_veri.values())).reindex(columns=tum_kolonlar)

    excel_ok = _excel_yaz_guvenli(df_genis, excel_dosya)
    hata_ok = hata_raporu_yaz(hatalar, hata_dosya)
    return excel_ok, hata_ok


# ============================================================
# CHECKPOINT'TEN EXCEL
# ============================================================

def jsondan_excel_uret(basla, son):
    yollar = dosya_yollari_uret(basla, son)
    veri = checkpoint_yukle(yollar["checkpoint"], basla, son, log_dosya=yollar["log"])
    if not veri:
        print("[UYARI] Bu tarih aralığına ait JSON ara kaydı bulunamadı.")
        return
    guncel = _lig_filtresi_imzasi()
    kayitli = veri.get("lig_filtresi")
    if kayitli and kayitli != guncel:
        print("[UYARI] Excel üretimi durduruldu: checkpoint aktif lig grubuyla eşleşmiyor.")
        print(f"  Checkpoint grubu: {kayitli}")
        print(f"  Aktif grup:       {guncel}")
        print("  Doğru --grup değerini kullanın veya ilgili grubun kayıt klasörünü seçin.")
        return
    genis_veri = {str(k): v for k, v in veri.get("genis_veri", {}).items()}
    hatalar = hata_listesini_dict_yap(veri.get("hatalar", {}))
    oran_kolonlari = []
    for kolon in veri.get("oran_kolonlari", []):
        oran_kolonu_ekle(oran_kolonlari, kolon)
    oran_kolonlarini_tamamla(oran_kolonlari, genis_veri)
    excel_ok, hata_ok = excel_ara_kaydet(
        genis_veri, oran_kolonlari, yollar["excel"], yollar["hata"], hatalar)
    if excel_ok:
        print(f"[TAMAM] {len(genis_veri):,} maç Excel'e aktarıldı.")
        print(f"Excel: {yollar['excel'].resolve()}")
    else:
        print("[HATA] Excel yazılamadı.")
    if hata_ok:
        print(f"Hata raporu: {yollar['hata'].resolve()}")
    else:
        print("[HATA] Hata raporu yazılamadı.")


# ============================================================
# TARİH
# ============================================================

def tarih_araligi_iste():
    while True:
        try:
            basla_girdi = input("Başlangıç tarihi (GG.AA.YYYY): ").strip()
            son_girdi = input("Bitiş tarihi (GG.AA.YYYY): ").strip()
            basla = datetime.strptime(basla_girdi, "%d.%m.%Y").strftime("%Y-%m-%d")
            son = datetime.strptime(son_girdi, "%d.%m.%Y").strftime("%Y-%m-%d")
            if basla > son:
                print("Hata: Başlangıç bitişten sonra olamaz.\n")
                continue
            return basla, son
        except ValueError:
            print("Tarih biçimi hatalı. Örnek: 01.01.2020\n")


# ============================================================
# DEVAM
# ============================================================

def checkpoint_mac_listesi_uret(genis_veri):
    satirlar = []
    for mac_id, bilgi in genis_veri.items():
        if not isinstance(bilgi, dict):
            continue
        mid = metin_temizle(bilgi.get("Maç ID") or mac_id)
        ev = metin_temizle(bilgi.get("Ev Sahibi", ""))
        dep = metin_temizle(bilgi.get("Deplasman", ""))
        if not mid or not ev or not dep:
            continue
        satirlar.append({
            "Tarih": metin_temizle(bilgi.get("Tarih", "")),
            "Saat": metin_temizle(bilgi.get("Saat", "")),
            "Ev Sahibi": ev, "Deplasman": dep,
            "Maç": metin_temizle(bilgi.get("Maç", "")) or f"{ev} vs {dep}",
            "İY Skor": _skor_duzenle(bilgi.get("İY Skor", "")),
            "Skor": _skor_duzenle(bilgi.get("Skor", "")),
            "Maç ID": mid,
            "İddaa Kodu": metin_temizle(bilgi.get("İddaa Kodu", "")),
            "Maç Slug": metin_temizle(bilgi.get("Maç Slug", "")),
        })
    if not satirlar:
        return pd.DataFrame()
    df = pd.DataFrame(satirlar)
    df["Maç ID"] = df["Maç ID"].astype(str)
    return df.drop_duplicates(subset=["Maç ID"], keep="first").reset_index(drop=True)


def mac_listelerini_birlestir(eski_df, yeni_df):
    alanlar = ["Tarih", "Saat", "Ev Sahibi", "Deplasman", "Maç",
               "İY Skor", "Skor", "Maç ID", "İddaa Kodu", "Maç Slug"]
    birlesik = {}
    for df in (eski_df, yeni_df):
        if df is None or df.empty:
            continue
        for kayit in df.to_dict(orient="records"):
            mid = metin_temizle(kayit.get("Maç ID", ""))
            if not mid:
                continue
            hedef = birlesik.setdefault(mid, {alan: "" for alan in alanlar})
            for alan in alanlar:
                deger = metin_temizle(kayit.get(alan, ""))
                if deger:
                    hedef[alan] = deger
            hedef["Maç ID"] = mid
    sonuc = pd.DataFrame(list(birlesik.values()), columns=alanlar)
    if not sonuc.empty:
        sonuc["Maç ID"] = sonuc["Maç ID"].astype(str)
    return sonuc



# ============================================================
# MAÇ LİSTESİ — vd.mackolik.com/livedata
# ============================================================

def iddaa_bilgi(basla, son, hedef_gunler=None, ara_kayit_dosyasi=None,
                 tamamlanan_gunler=None):
    """Günlük futbol maçlarını çeker; her günün ara kaydını atomik yazar."""
    baslangic = datetime.strptime(basla, "%Y-%m-%d").date()
    bitis = datetime.strptime(son, "%Y-%m-%d").date()
    if baslangic > bitis:
        raise ValueError("Başlangıç tarihi bitiş tarihinden büyük olamaz.")

    tamamlanan = set(str(g) for g in (tamamlanan_gunler or set()))
    hedefler = (list(dict.fromkeys(str(g) for g in hedef_gunler))
                if hedef_gunler is not None else _gun_araligi(basla, son))
    hedefler = [g for g in hedefler if g not in tamamlanan]
    maclar, hatalar = [], []
    api_url = "https://vd.mackolik.com/livedata"

    if ara_kayit_dosyasi:
        try:
            ara_yol = Path(ara_kayit_dosyasi)
            if ara_yol.exists():
                with open(ara_yol, "r", encoding="utf-8") as f:
                    ara = json.load(f)
                if (isinstance(ara, dict)
                        and ara.get("baslangic_tarihi") == basla
                        and ara.get("bitis_tarihi") == son):
                    eski_maclar = ara.get("maclar", [])
                    if isinstance(eski_maclar, list):
                        maclar.extend(k for k in eski_maclar if isinstance(k, dict))
                    ara_gunler = ara.get("tamamlanan_gunler", [])
                    if isinstance(ara_gunler, list):
                        tamamlanan.update(str(g) for g in ara_gunler)
                    hedefler = [g for g in hedefler if g not in tamamlanan]
        except (OSError, ValueError, TypeError) as hata:
            print(f"[UYARI] Maç listesi ara kaydı okunamadı; yeniden alınacak: {hata}")

    def _ara_kaydi_yaz():
        if ara_kayit_dosyasi:
            _guvenli_json_yaz(Path(ara_kayit_dosyasi), {
                "baslangic_tarihi": basla, "bitis_tarihi": son,
                "maclar": maclar, "tamamlanan_gunler": sorted(tamamlanan),
                "guncelleme": datetime.now().astimezone().isoformat(timespec="seconds"),
            })

    for tarih_iso in hedefler:
        gun = datetime.strptime(tarih_iso, "%Y-%m-%d").date()
        print(f"[MAÇ LİSTESİ] {tarih_iso}")
        try:
            cevap = istek_yap(api_url, params={"date": gun.strftime("%d/%m/%Y")},
                              livedata=True)
            veri = cevap.json()
            if not isinstance(veri, dict):
                raise ValueError("livedata yanıtı JSON nesnesi değil.")
            kayitlar = veri.get("m") or []
            if not isinstance(kayitlar, list):
                raise ValueError("livedata 'm' alanı liste değil.")
        except ErisimKisitlamaHatasi:
            _ara_kaydi_yaz()
            raise
        except Exception as hata:
            hatalar.append({"Tarih": tarih_iso, "ID": "", "Maç": "",
                            "Hata": f"LIVEDATA: maç listesi çekilemedi: {hata}"})
            print(f"[HATA] {tarih_iso}: {hata}")
            _ara_kaydi_yaz()
            continue

        gun_mac_sayisi = 0
        for kayit in kayitlar:
            try:
                if not isinstance(kayit, list) or len(kayit) < 37:
                    continue
                yarisma = kayit[36]
                if not isinstance(yarisma, list) or len(yarisma) < 12:
                    continue
                if sayiya_cevir(yarisma[11]) != 1 or not _lig_uygun_mu(yarisma):
                    continue
                mac_id = metin_temizle(kayit[0])
                ev, dep = metin_temizle(kayit[2]), metin_temizle(kayit[4])
                if not mac_id.isdigit() or not ev or not dep:
                    continue
                # İddaa kodu boş olan futbol maçları da korunur.
                iddaa_kodu = metin_temizle(kayit[14]) if len(kayit) > 14 else ""
                iy_skor = metin_temizle(kayit[7]) if len(kayit) > 7 else ""
                if not iy_skor:
                    iy_ev = sayiya_cevir(kayit[10]) if len(kayit) > 10 else None
                    iy_dep = sayiya_cevir(kayit[11]) if len(kayit) > 11 else None
                    if iy_ev is not None and iy_dep is not None:
                        iy_skor = f"{iy_ev}-{iy_dep}"
                ev_gol = sayiya_cevir(kayit[12]) if len(kayit) > 12 else None
                dep_gol = sayiya_cevir(kayit[13]) if len(kayit) > 13 else None
                skor = f"{ev_gol}-{dep_gol}" if ev_gol is not None and dep_gol is not None else ""
                saat = metin_temizle(kayit[16]) if len(kayit) > 16 else ""
                tarih_kayit = metin_temizle(kayit[35]) if len(kayit) > 35 else ""
                tarih_norm = gun.strftime("%d-%m-%Y")
                if tarih_kayit:
                    tarih_kayit = tarih_kayit.replace("/", "-").replace(".", "-")
                    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d-%m-%y"):
                        try:
                            tarih_norm = datetime.strptime(tarih_kayit, fmt).strftime("%d-%m-%Y")
                            break
                        except ValueError:
                            continue
                maclar.append({
                    "Tarih": tarih_norm, "Saat": saat, "Ev Sahibi": ev,
                    "Deplasman": dep, "Maç": f"{ev} vs {dep}",
                    "İY Skor": iy_skor, "Skor": skor, "Maç ID": mac_id,
                    "İddaa Kodu": iddaa_kodu, "Maç Slug": "",
                })
                gun_mac_sayisi += 1
            except (IndexError, TypeError, ValueError):
                continue
        tamamlanan.add(tarih_iso)
        print(f"  → {gun_mac_sayisi} futbol maçı")
        maclar = list({str(k.get("Maç ID", "")): k for k in maclar
                       if metin_temizle(k.get("Maç ID", ""))}.values())
        _ara_kaydi_yaz()

    df = pd.DataFrame(maclar, columns=[
        "Tarih", "Saat", "Ev Sahibi", "Deplasman", "Maç", "İY Skor",
        "Skor", "Maç ID", "İddaa Kodu", "Maç Slug",
    ])
    if not df.empty:
        df["Maç ID"] = df["Maç ID"].astype(str)
        df = df.drop_duplicates(subset=["Maç ID"], keep="first").reset_index(drop=True)
    return df, hatalar


# ═════════════════════════════════════════════════════════════
# v15.7 DÜZELTME #2:
#   _eksik_gunleri_bul artık df_mac'e BAKMIYOR. Yalnızca
#   tamamlanan_gunler güvenilir. Gün başarıyla tamamlandığında
#   iddaa_bilgi hatasızdönersin; _tamamlanan_gunleri_guncelle
#   ekler. Aksi hâlde yeniden denenir.
# ═════════════════════════════════════════════════════════════
def _eksik_gunleri_cevir(df_mac, hatalar, basla, son, ara_liste_dosyasi,
                          tamamlanan_gunler=None, onceki_durum=None):
    if tamamlanan_gunler is None:
        tamamlanan_gunler = set()
    eksik_gunler = _eksik_gunleri_bul(basla, son, tamamlanan_gunler)
    if not eksik_gunler:
        if onceki_durum and "Kısmi" in onceki_durum:
            return df_mac, hatalar, onceki_durum
        return df_mac, hatalar, "Tamam"
    print(f"[DEVAM] {len(eksik_gunler)} gün eksik; livedata çağrılıyor...")
    eksik_df, eksik_hatalari = iddaa_bilgi(
        basla, son, hedef_gunler=eksik_gunler,
        ara_kayit_dosyasi=ara_liste_dosyasi,
        tamamlanan_gunler=tamamlanan_gunler)

    # v15.7: hatası olmayan günler tamamlandı olarak işaretlenir
    _tamamlanan_gunleri_guncelle(eksik_hatalari, eksik_gunler, tamamlanan_gunler)

    tarih_hatalarini_temizle(hatalar, eksik_gunler)
    hatalari_birlestir(hatalar, eksik_hatalari)
    df_mac = mac_listelerini_birlestir(df_mac, eksik_df)
    durum = "Tamam" if not eksik_hatalari else "Kısmi"
    return df_mac, hatalar, durum


def _tamamlanma_kriteri(bilgi):
    if not isinstance(bilgi, dict):
        return False
    return (bilgi.get("Arşiv Sayfa Durumu") == "Tamam"
            and bilgi.get("Oran Durumu") == "Oran bulundu")


def _kadro_eksik_mi(bilgi):
    if not MATCHDATA_AKTIF:
        return False
    if not _tamamlanma_kriteri(bilgi):
        return False
    ev_eksik = not bool(metin_temizle(bilgi.get("Ev Kadro", "")))
    dep_eksik = not bool(metin_temizle(bilgi.get("Deplasman Kadro", "")))
    return ev_eksik or dep_eksik


def _olaylar_bos_mu(bilgi):
    alanlar = ["Goller", "Sarı Kartlar", "Kırmızı Kartlar",
               "Değişiklikler", "Kaçan Penaltılar"]
    return not any(metin_temizle(bilgi.get(k, "")) for k in alanlar)


def _matchdata_denenmeli_mi(bilgi):
    if not MATCHDATA_AKTIF:
        return False
    if not isinstance(bilgi, dict):
        return False
    # v15.7: yalnızca başarılı MatchData çağrısı işaretlenir
    if bilgi.get("__matchdata_denendi") == "1":
        return False
    if not metin_temizle(bilgi.get("Ev Kadro", "")):
        return True
    if not metin_temizle(bilgi.get("Deplasman Kadro", "")):
        return True
    if _olaylar_bos_mu(bilgi):
        return True
    return False


# ============================================================
# TEK MAÇ
# ============================================================

# ═════════════════════════════════════════════════════════════
# v15.7 DÜZELTME #5:
#   __matchdata_denendi artık YALNIZCA MatchData başarıyla
#   döndüğünde işaretleniyor. Transient hata olursa sonraki
#   turda tekrar denenir (MAKS_TUR'a kadar).
# ═════════════════════════════════════════════════════════════
async def _tek_mac_isle_async(row, tur, genis_veri, hatalar, oran_kolonlari):
    mac_id = str(row["Maç ID"])
    bilgi = genis_veri[mac_id]

    if (_tamamlanma_kriteri(bilgi) and _matchdata_denenmeli_mi(bilgi)):
        print(f"[TUR {tur}] [FALLBACK] {mac_id} - Sadece kadro/olay eksik, "
              f"MatchData çekiliyor...")
        bilgi, hatalar, basarili = await _matchdata_uygula_async(
            bilgi, mac_id, hatalar)
        if basarili:
            bilgi["__matchdata_denendi"] = "1"
        else:
            print(f"[TUR {tur}] [FALLBACK] {mac_id} - MatchData başarısız, "
                  f"sonraki turda tekrar denenecek.")
        return ("ok" if basarili else "matchdata_eksik"), hatalar

    try:
        sonuc = await get_bahis_oranlari_async(row.to_dict())

        for alan, deger in sonuc["mac_bilgisi"].items():
            if deger not in (None, ""):
                bilgi[alan] = deger
        bilgi["Kaynak URL"] = sonuc["kaynak_url"]
        _bilgiden_oran_istatistik_sil(bilgi, oran_kolonlari)

        gecerli_oranlar = []
        for bahis in sonuc["oranlar"]:
            market = metin_temizle(bahis.get("Bahis Türü", ""))
            secim = metin_temizle(bahis.get("Seçim", ""))
            oran = oran_temizle(bahis.get("Oran", ""))
            if not oran:
                continue
            kolon = bahis_basligi(market, secim)
            oran_kolonu_ekle(oran_kolonlari, kolon)
            gecerli_oranlar.append((kolon, oran))

        if gecerli_oranlar:
            for kolon, oran in gecerli_oranlar:
                bilgi[kolon] = oran
            bilgi["Oran Durumu"] = "Oran bulundu"
            bilgi["Arşiv Sayfa Durumu"] = "Tamam"
            hatalar = hata_kayitlarini_temizle(
                hatalar, mac_id,
                icerenler=("oran bulunamadı",),
                baslangiclar=("ARŞİV: HTTP",))
        else:
            bilgi["Oran Durumu"] = "Arşiv sayfasında oran bulunamadı"
            bilgi["Arşiv Sayfa Durumu"] = "Oran Eksik"
            hatalar = hata_kayitlarini_temizle(
                hatalar, mac_id, icerenler=("oran bulunamadı",))
            hata_ekle(hatalar, {
                "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
                "Maç": bilgi.get("Maç", ""),
                "Hata": f"Arşiv sayfası açıldı ancak geçerli oran yok. "
                        f"Tur {tur}/{MAKS_TUR}."
            })

        for kolon, deger in sonuc["istatistikler"].items():
            bilgi[kolon] = deger
        ist_durum = sonuc.get("istatistik_durum", "kapsayici_yok")
        if str(ist_durum).startswith("bulundu"):
            bilgi["İstatistik Durumu"] = ist_durum
        elif ist_durum == "kapsayici_yok":
            bilgi["İstatistik Durumu"] = "İstatistik kapsayıcısı yok"
        elif str(ist_durum).startswith("hata"):
            bilgi["İstatistik Durumu"] = ist_durum
            hata_ekle(hatalar, {
                "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
                "Maç": bilgi.get("Maç", ""),
                "Hata": f"İSTATİSTİK: {ist_durum}",
            })
        else:
            bilgi["İstatistik Durumu"] = "İstatistik bilinmiyor"

        if not bilgi.get("İY Skor"):
            bilgi["İY Skor Kaynağı"] = ""
            hatalar = hata_kayitlarini_temizle(
                hatalar, mac_id, baslangiclar=("İY SKOR:",))
            hata_ekle(hatalar, {
                "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
                "Maç": bilgi.get("Maç", ""),
                "Hata": "İY SKOR: arşiv alanı/JavaScript ve maç listesi "
                        "kaynağında skor bulunamadı."
            })
        else:
            hatalar = hata_kayitlarini_temizle(
                hatalar, mac_id, baslangiclar=("İY SKOR:",))
        hatalar = hata_kayitlarini_temizle(
            hatalar, mac_id, baslangiclar=("ARŞİV: HTTP",))

        durum_sonuc = "ok"
        if MATCHDATA_AKTIF and _matchdata_denenmeli_mi(bilgi):
            etiket = ("kadro" if (not metin_temizle(bilgi.get("Ev Kadro", ""))
                                   or not metin_temizle(bilgi.get("Deplasman Kadro", "")))
                      else "olaylar")
            print(f"[TUR {tur}] [FALLBACK] {mac_id} - {etiket} eksik, "
                  f"MatchData çekiliyor...")
            bilgi, hatalar, basarili = await _matchdata_uygula_async(
                bilgi, mac_id, hatalar)
            if basarili:
                bilgi["__matchdata_denendi"] = "1"
            else:
                # Arşiv oranı başarılı olsa da bu turdaki ek MatchData işlemi
                # başarısızdır; tur istatistiklerinde ayrıca görünür.
                durum_sonuc = "matchdata_eksik"
                print(f"[TUR {tur}] [FALLBACK] {mac_id} - MatchData başarısız; "
                      "sonraki turda tekrar denenecek.")

        return durum_sonuc, hatalar

    except ErisimKisitlamaHatasi as hata:
        bilgi["Arşiv Sayfa Durumu"] = f"Erişim Hatası (HTTP {hata.durum_kodu})"
        bilgi["Oran Durumu"] = f"Çekilemedi (HTTP {hata.durum_kodu})"
        bilgi["İstatistik Durumu"] = f"Çekilemedi (HTTP {hata.durum_kodu})"
        hatalar = hata_kayitlarini_temizle(
            hatalar, mac_id, baslangiclar=("ARŞİV:",))
        hata_ekle(hatalar, {
            "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
            "Maç": bilgi.get("Maç", ""),
            "Hata": f"ARŞİV: HTTP {hata.durum_kodu}; URL={hata.url}"
        })
        if hata.durum_kodu == 429:
            print(f"[TUR {tur}] [MAÇ HTTP 429] {mac_id} | atlandı.")
            return "kisitlama_429", hatalar
        elif hata.durum_kodu == 401:
            print(f"[TUR {tur}] [MAÇ HTTP 401] {mac_id} | kalıcı erişim reddi.")
            return "kisitlama_401", hatalar
        else:
            print(f"[TUR {tur}] [MAÇ HTTP {hata.durum_kodu}] {mac_id} | "
                  f"kalıcı erişim reddi.")
            return "kisitlama_403", hatalar
    except Exception as hata:
        bilgi["Arşiv Sayfa Durumu"] = "Hata"
        bilgi["Oran Durumu"] = "Çekilemedi (parse/ağ hatası)"
        bilgi["İstatistik Durumu"] = "Çekilemedi (parse/ağ hatası)"
        hatalar = hata_kayitlarini_temizle(
            hatalar, mac_id, baslangiclar=("ARŞİV:",))
        hata_ekle(hatalar, {
            "Tarih": bilgi.get("Tarih", ""), "ID": mac_id,
            "Maç": bilgi.get("Maç", ""), "Hata": f"ARŞİV: {hata}"
        })
        return "hata", hatalar


# ============================================================
# ANA DÖNGÜ
# ============================================================

async def iddaa_bahis_oranlari_async(basla, son):
    oturum_baslangici = time.time()
    yollar = dosya_yollari_uret(basla, son)
    excel_dosya = yollar["excel"]
    hata_dosya = yollar["hata"]
    ara_liste_dosyasi = yollar["mac_listesi_gecici"]

    mac_log = CheckpointLog(yollar["log"])
    kayit_locksu = asyncio.Lock()

    genis_veri = {}
    hatalar = {}
    oran_kolonlari = []
    tamamlanan_gunler = set()
    df_mac = pd.DataFrame()
    df_mac_records = []
    mac_listesi_durum = None
    toplam_mac = 0
    df_mac_ids = set()
    temp_silindi = False
    kalici_hata_idler = set()

    print("=" * 76)
    print(f"Maçkolik ARŞİV v15.11: {basla} → {son}")
    print(f"Grup: {AKTIF_ETIKET}  ({', '.join(AKTIF_GRUPLAR)})")
    print(f"Eşzamanlı: {ESZAMANLI_ISTEK} | Min istek: {MIN_ISTEK_ARASI}s | "
          f"Jitter: {ISTEK_JITTER_ASYNC}s")
    print(f"TLS Impersonate: {CURL_IMPERSONATE}")
    print(f"Kayıt klasörü: {yollar['klasor'].resolve()}")
    print("=" * 76)

    async def guvenli_tam_kaydet(
        yollar_, basla_, son_, genis_veri_, hatalar_, oran_kolonlari_,
        mac_listesi_, mac_listesi_durum_, excel_dosya_, hata_dosya_,
        mac_log_, excel_yazilsin_=True, tamamlanan_gunler_=None,
        kalici_hata_idler_=None
    ):
        """JSON kilitli, Excel snapshot + thread."""
        async with kayit_locksu:
            try:
                checkpoint_yaz(
                    yollar_, basla_, son_, genis_veri_, hatalar_,
                    oran_kolonlari_,
                    mac_listesi=mac_listesi_,
                    mac_listesi_durum=mac_listesi_durum_,
                    tamamlanan_gunler=tamamlanan_gunler_,
                    kalici_hata_idler=kalici_hata_idler_,
                )
            except Exception as e:
                print(f"[KAYIT HATASI] JSON: {e}")
                return False, False, False

            if mac_log_ is not None:
                try:
                    mac_log_.temizle()
                except Exception:
                    pass

            if not excel_yazilsin_:
                return True, False, False

            excel_snap = {
                k: dict(v) if isinstance(v, dict) else v
                for k, v in genis_veri_.items()
            }
            hata_snap = {k: list(v) for k, v in hatalar_.items()
                         if isinstance(v, list)}
            oran_snap = list(oran_kolonlari_)

            try:
                excel_ok, hata_ok = await asyncio.to_thread(
                    excel_ara_kaydet, excel_snap, oran_snap,
                    excel_dosya_, hata_dosya_, hata_snap)
            except Exception as e:
                print(f"[KAYIT HATASI] Excel: {e}")
                return True, False, False

            if not excel_ok:
                print("[KAYIT] Excel yazılamadı (JSON korundu).")
            if not hata_ok:
                print("[KAYIT] Hata raporu yazılamadı (JSON korundu).")
            return True, excel_ok, hata_ok

    kayit = None
    try:
        kayit = checkpoint_yukle(yollar["checkpoint"], basla, son,
                                  log_dosya=yollar["log"])

        # ═════════════════════════════════════════════════════
        # v15.7 DÜZELTME #1:
        #   Lig imzası değişmişse checkpoint TAMAMEN reddedilir.
        #   genis_veri, hatalar, tamamlanan_gunler, kalici_hata_idler
        #   boş kalır; maç listesi yeniden çekilir; Excel yeni
        #   gruba ait verilerle yazılır.
        # ═════════════════════════════════════════════════════
        guncel_filtre = _lig_filtresi_imzasi()
        if (kayit and kayit.get("lig_filtresi")
                and kayit.get("lig_filtresi") != guncel_filtre):
            print("[UYARI] Lig filtresi değişmiş; checkpoint YÜKLENMİYOR.")
            print(f"  Eski: {kayit.get('lig_filtresi')}")
            print(f"  Yeni: {guncel_filtre}")
            print("  Eski checkpoint dosyası diskte kalır, ilk kayıtta üzerine yazılır.")
            kayit = None

        if kayit:
            genis_veri = {str(k): v for k, v in kayit.get("genis_veri", {}).items()}
            hatalar = hata_listesini_dict_yap(kayit.get("hatalar", {}))
            tamamlanan_gunler = set(kayit.get("tamamlanan_gunler", set()))
            kalici_hata_idler = set(kayit.get("kalici_hata_idler", []) or [])
            for kolon in kayit.get("oran_kolonlari", []):
                oran_kolonu_ekle(oran_kolonlari, kolon)
            oran_kolonlarini_tamamla(oran_kolonlari, genis_veri)

        # ─── Maç listesi ─────────────────────────────────────
        if (kayit and isinstance(kayit.get("mac_listesi"), list)
                and kayit["mac_listesi"]):
            df_mac = pd.DataFrame(kayit["mac_listesi"])
            mac_listesi_durum = kayit.get("mac_listesi_durum") or "Tamam"
            print(f"[DEVAM] Maç listesi checkpoint'ten yüklendi: {len(df_mac):,}")
            df_mac, hatalar, mac_listesi_durum = _eksik_gunleri_cevir(
                df_mac, hatalar, basla, son, ara_liste_dosyasi,
                tamamlanan_gunler, onceki_durum=mac_listesi_durum)
        elif kayit:
            eski_mac_df = checkpoint_mac_listesi_uret(genis_veri)
            if not eski_mac_df.empty:
                df_mac = eski_mac_df.copy()
                print(f"[DEVAM] Checkpoint'ten maç listesi yeniden "
                      f"kuruldu: {len(df_mac):,}")
                df_mac, hatalar, mac_listesi_durum = _eksik_gunleri_cevir(
                    df_mac, hatalar, basla, son, ara_liste_dosyasi,
                    tamamlanan_gunler,
                    onceki_durum=kayit.get("mac_listesi_durum"))
            else:
                print("[MAÇ LİSTESİ] Checkpoint'te maç listesi yok; yeniden çekiliyor...")
                df_mac, liste_hatalari = iddaa_bilgi(
                    basla, son, ara_kayit_dosyasi=ara_liste_dosyasi,
                    tamamlanan_gunler=tamamlanan_gunler)
                _tamamlanan_gunleri_guncelle(
                    liste_hatalari, _gun_araligi(basla, son), tamamlanan_gunler)
                hatalari_birlestir(hatalar, liste_hatalari)
                mac_listesi_durum = "Tamam" if not liste_hatalari else "Kısmi"
        else:
            print("[MAÇ LİSTESİ] İlk çalıştırma; maç listesi alınıyor...")
            df_mac, liste_hatalari = iddaa_bilgi(
                basla, son, ara_kayit_dosyasi=ara_liste_dosyasi,
                tamamlanan_gunler=tamamlanan_gunler)
            _tamamlanan_gunleri_guncelle(
                liste_hatalari, _gun_araligi(basla, son), tamamlanan_gunler)
            hatalari_birlestir(hatalar, liste_hatalari)
            mac_listesi_durum = "Tamam" if not liste_hatalari else "Kısmi"

        hatalar = hatalari_tekrarsiz_yap(hatalar)

        if df_mac.empty:
            print("Maç listesi boş veya alınamadı.")
            return

        df_mac["Maç ID"] = df_mac["Maç ID"].astype(str)
        df_mac = df_mac.drop_duplicates(subset=["Maç ID"], keep="first").reset_index(drop=True)
        toplam_mac = len(df_mac)
        df_mac_ids = set(df_mac["Maç ID"].astype(str).tolist())
        df_mac_records = df_mac.to_dict(orient="records")

        for _, row in df_mac.iterrows():
            mac_id = str(row["Maç ID"])
            if mac_id not in genis_veri:
                genis_veri[mac_id] = {kolon: "" for kolon in MAC_KOLONLARI}
            bilgi = genis_veri[mac_id]
            for alan in ["Tarih", "Saat", "Ev Sahibi", "Deplasman",
                         "Maç", "Skor", "İddaa Kodu"]:
                yeni = metin_temizle(row.get(alan, ""))
                if yeni and (not bilgi.get(alan) or alan in {"İddaa Kodu"}):
                    bilgi[alan] = yeni
            bilgi["Maç ID"] = mac_id
            if not bilgi.get("İY Skor") and metin_temizle(row.get("İY Skor", "")):
                bilgi["İY Skor"] = _skor_duzenle(row.get("İY Skor", ""))
                bilgi["İY Skor Kaynağı"] = "Maç listesi ilk yarı skoru"
            elif bilgi.get("İY Skor") and not bilgi.get("İY Skor Kaynağı"):
                bilgi["İY Skor Kaynağı"] = "Önceki kayıt; kaynak bilgisi yok"
            bilgi.setdefault("Arşiv Sayfa Durumu", "")
            bilgi.setdefault("Oran Durumu", "Oran bulunamadı")
            bilgi.setdefault("İstatistik Durumu", "İstatistik bulunamadı")

        oran_kolonlarini_tamamla(oran_kolonlari, genis_veri)
        tamamlanan_idler = {
            mid for mid in df_mac_ids
            if _tamamlanma_kriteri(genis_veri.get(mid, {}))
        }
        kadro_eksik_idler = {
            mid for mid in df_mac_ids if _kadro_eksik_mi(genis_veri.get(mid, {}))
        }
        print(f"Toplam maç: {toplam_mac:,}")
        print(f"Önceden tamamlanan (sayfa+oran): {len(tamamlanan_idler):,}")
        print(f"Tamamlanan gün sayısı: {len(tamamlanan_gunler):,}")
        if MATCHDATA_AKTIF:
            print(f"Kadro eksiği olan: {len(kadro_eksik_idler):,}")
        if kalici_hata_idler:
            print(f"Kalıcı 401/403 (atlanacak): {len(kalici_hata_idler):,}")

        oturumda_islenen = 0
        son_kayit = 0
        hala_basarisiz_idler = []

        for tur in range(1, MAKS_TUR + 1):
            kalan_idler_tur = [
                mid for mid in df_mac["Maç ID"].astype(str).tolist()
                if mid not in kalici_hata_idler
                and (not _tamamlanma_kriteri(genis_veri.get(mid, {}))
                     or _matchdata_denenmeli_mi(genis_veri.get(mid, {})))
            ]
            if not kalan_idler_tur:
                print(f"\n[TUR {tur}/{MAKS_TUR}] Tekrar denenecek maç yok. "
                      f"Tur bitti.")
                break

            kalan_df_tur = df_mac[
                df_mac["Maç ID"].astype(str).isin(kalan_idler_tur)
            ].copy()
            print(f"\n[TUR {tur}/{MAKS_TUR}] Bu turda denenecek maç: "
                  f"{len(kalan_df_tur):,}")
            if kalici_hata_idler:
                print(f"  ({len(kalici_hata_idler):,} maç kalıcı 401/403, atlanıyor.)")

            tur_baslangic_tamamlanan = len(tamamlanan_idler)
            tur_kisitlama_429 = 0
            tur_kisitlama_401 = 0
            tur_kisitlama_403 = 0
            tur_hata = 0
            tur_matchdata_eksik = 0

            semaphore = asyncio.Semaphore(ESZAMANLI_ISTEK)

            # ═════════════════════════════════════════════════
            # v15.7 DÜZELTME #6:
            #   process_row'un TAMAMI korumalı; gather sızıntısı yok.
            #   Her aşama kendi try/except'inde. Beklenmeyen bir
            #   istisna olsa bile maç ID'siyle kaydedilir.
            # ═════════════════════════════════════════════════
            async def process_row(row, tur=tur):
                nonlocal oturumda_islenen, son_kayit, temp_silindi
                nonlocal tur_kisitlama_429, tur_kisitlama_401
                nonlocal tur_kisitlama_403, tur_hata, tur_matchdata_eksik

                mac_id = str(row["Maç ID"])
                durum = "hata"
                hatalar_guncel = hatalar

                # 1. Asıl işlem
                try:
                    async with semaphore:
                        durum, hatalar_guncel = await _tek_mac_isle_async(
                            row, tur, genis_veri, hatalar, oran_kolonlari)
                except Exception as exc:
                    import traceback
                    traceback.print_exc()
                    try:
                        hata_ekle(hatalar, {
                            "Tarih": genis_veri.get(mac_id, {}).get("Tarih", ""),
                            "ID": mac_id,
                            "Maç": genis_veri.get(mac_id, {}).get("Maç", ""),
                            "Hata": f"BEKLENMEYEN: {type(exc).__name__}: {exc}",
                        })
                    except Exception:
                        pass
                    if mac_id in genis_veri:
                        genis_veri[mac_id]["Arşiv Sayfa Durumu"] = "Hata (beklenmeyen)"
                    durum, hatalar_guncel = "hata", hatalar

                # 2. Log
                try:
                    mac_log.mac_yaz(
                        mac_id, genis_veri.get(mac_id, {}),
                        hatalar_guncel.get(f"id:{mac_id}") if isinstance(
                            hatalar_guncel, dict) else None)
                except Exception as log_hatasi:
                    print(f"[LOG HATASI] {mac_id}: {log_hatasi}")

                # 3. Sayaçlar
                oturumda_islenen += 1
                if durum in ("kisitlama_401", "kisitlama_403"):
                    kalici_hata_idler.add(mac_id)

                # 4. Kayıt
                kritik = durum in ("kisitlama_429", "kisitlama_401",
                                    "kisitlama_403")
                esik_asildi = ((oturumda_islenen - son_kayit)
                               >= EXCEL_KAYIT_ARALIGI)
                if kritik or esik_asildi:
                    try:
                        json_ok, excel_ok, hata_ok = await guvenli_tam_kaydet(
                            yollar, basla, son, genis_veri, hatalar,
                            oran_kolonlari,
                            mac_listesi_=df_mac_records,
                            mac_listesi_durum_=mac_listesi_durum,
                            excel_dosya_=excel_dosya, hata_dosya_=hata_dosya,
                            mac_log_=mac_log, excel_yazilsin_=True,
                            tamamlanan_gunler_=tamamlanan_gunler,
                            kalici_hata_idler_=kalici_hata_idler,
                        )
                    except Exception as kayit_hatasi:
                        print(f"[KAYIT HATASI] guvenli_tam_kaydet: "
                              f"{kayit_hatasi}")
                        json_ok = excel_ok = hata_ok = False

                    if json_ok:
                        son_kayit = oturumda_islenen
                        if excel_ok and hata_ok:
                            print(f"[TUR {tur}] [KAYIT] Checkpoint + Excel + "
                                  f"Hata raporu güncellendi.")
                        else:
                            parcalar = ["Checkpoint OK"]
                            parcalar.append("Excel OK" if excel_ok
                                            else "Excel BAŞARISIZ")
                            parcalar.append("Hata raporu OK" if hata_ok
                                            else "Hata raporu BAŞARISIZ")
                            print(f"[TUR {tur}] [KAYIT] " + " | ".join(parcalar))

                        if (not temp_silindi
                                and ara_liste_dosyasi.exists()
                                and df_mac_records
                                and _checkpoint_mac_listesi_dolu(yollar)):
                            try:
                                ara_liste_dosyasi.unlink()
                                temp_silindi = True
                                print("[MAÇ LİSTESİ] Geçici ara kayıt "
                                      "checkpoint doğrulandıktan sonra silindi.")
                            except Exception:
                                pass
                    else:
                        print(f"[TUR {tur}] [KAYIT HATASI] Checkpoint "
                              f"yazılamadı; eşik ilerletilmiyor.")

                # 5. Durum güncelle
                try:
                    if _tamamlanma_kriteri(genis_veri.get(mac_id, {})):
                        tamamlanan_idler.add(mac_id)
                except Exception:
                    pass

                if durum == "kisitlama_429":
                    tur_kisitlama_429 += 1
                elif durum == "kisitlama_401":
                    tur_kisitlama_401 += 1
                elif durum == "kisitlama_403":
                    tur_kisitlama_403 += 1
                elif durum == "matchdata_eksik":
                    tur_matchdata_eksik += 1
                elif durum == "hata":
                    tur_hata += 1

                # 6. Konsol
                try:
                    gecen = time.time() - oturum_baslangici
                    bilgi = genis_veri.get(mac_id, {})
                    ev_k = "E:✓" if metin_temizle(bilgi.get("Ev Kadro", "")) else "E:✗"
                    dep_k = "D:✓" if metin_temizle(bilgi.get("Deplasman Kadro", "")) else "D:✗"
                    ist_flag = ("İ:✓" if str(bilgi.get("İstatistik Durumu", ""))
                                .startswith("bulundu") else "İ:✗")
                    print(
                        f"[TUR {tur}] İşlem: {oturumda_islenen:,} | "
                        f"Tamamlanan: {len(tamamlanan_idler):,}/{toplam_mac:,} | "
                        f"Sayfa: {bilgi.get('Arşiv Sayfa Durumu', '')} | "
                        f"Oran: {bilgi.get('Oran Durumu', '')} | "
                        f"{ev_k}/{dep_k} | {ist_flag} | "
                        f"Geçen: {sure_formatla(gecen)}"
                    )
                except Exception:
                    pass

            tasks = [process_row(row, tur) for _, row in kalan_df_tur.iterrows()]
            results = await asyncio.gather(*tasks, return_exceptions=True)

            # Yedek güvenlik ağı (process_row artık hep kendi içinde yakalıyor
            # ama yine de son savunma)
            for i, r in enumerate(results):
                if isinstance(r, Exception):
                    try:
                        row = kalan_df_tur.iloc[i]
                        mid = str(row["Maç ID"])
                    except Exception:
                        continue
                    try:
                        hata_ekle(hatalar, {
                            "Tarih": genis_veri.get(mid, {}).get("Tarih", ""),
                            "ID": mid,
                            "Maç": genis_veri.get(mid, {}).get("Maç", ""),
                            "Hata": f"GATHER: {type(r).__name__}: {r}",
                        })
                    except Exception:
                        pass
                    if mid in genis_veri:
                        genis_veri[mid]["Arşiv Sayfa Durumu"] = "Hata (gather)"

            hala_basarisiz_idler = [
                mid for mid in kalan_idler_tur
                if mid not in kalici_hata_idler
                and (not _tamamlanma_kriteri(genis_veri.get(mid, {}))
                     or _matchdata_denenmeli_mi(genis_veri.get(mid, {})))
            ]
            hala_basarisiz = len(hala_basarisiz_idler)
            bu_tur_basarili = len(kalan_df_tur) - hala_basarisiz

            print(
                f"\n[TUR {tur}/{MAKS_TUR}] Bitti. "
                f"Bu turda tamamlanan: {bu_tur_basarili}/{len(kalan_df_tur)} | "
                f"429: {tur_kisitlama_429} | "
                f"401: {tur_kisitlama_401} | 403: {tur_kisitlama_403} | "
                f"MatchData eksik: {tur_matchdata_eksik} | "
                f"diğer: {tur_hata} | hâlâ başarısız: {hala_basarisiz}"
            )

            try:
                json_ok, excel_ok, hata_ok = await guvenli_tam_kaydet(
                    yollar, basla, son, genis_veri, hatalar, oran_kolonlari,
                    mac_listesi_=df_mac_records,
                    mac_listesi_durum_=mac_listesi_durum,
                    excel_dosya_=excel_dosya, hata_dosya_=hata_dosya,
                    mac_log_=mac_log, excel_yazilsin_=True,
                    tamamlanan_gunler_=tamamlanan_gunler,
                    kalici_hata_idler_=kalici_hata_idler,
                )
                if json_ok:
                    print(f"[TUR {tur}] [KAYIT] Tur sonu JSON yazıldı.")
                    if not excel_ok:
                        print(f"[TUR {tur}] [KAYIT] Excel BAŞARISIZ.")
                    if not hata_ok:
                        print(f"[TUR {tur}] [KAYIT] Hata raporu BAŞARISIZ.")
                else:
                    print(f"[TUR {tur}] [KAYIT HATASI] Tur sonu JSON yazılamadı!")
            except Exception as kayit_hatasi:
                print(f"[TUR {tur}] [KAYIT HATASI] {kayit_hatasi}")

            if hala_basarisiz == 0:
                print(f"[TUR {tur}] Tüm maçlar başarıyla tamamlandı.")
                break
            if tur < MAKS_TUR and hala_basarisiz > 0:
                print(f"[TUR {tur}] {hala_basarisiz} maç başarısız; "
                      f"sonraki turda yeniden denenecek.")
        else:
            print(f"\n[TUR DÖNGÜSÜ] {MAKS_TUR} tur tamamlandı.")
            if hala_basarisiz_idler:
                print(f"[KALAN] {len(hala_basarisiz_idler)} maç hâlâ eksik:")
                for mid in hala_basarisiz_idler[:20]:
                    b = genis_veri.get(mid, {})
                    ev = "✗" if not metin_temizle(b.get("Ev Kadro", "")) else "✓"
                    dep = "✗" if not metin_temizle(b.get("Deplasman Kadro", "")) else "✓"
                    print(f"  - {mid} | {b.get('Maç', '')} | Ev:{ev} Dep:{dep}")
                if len(hala_basarisiz_idler) > 20:
                    print(f"  ... ve {len(hala_basarisiz_idler) - 20} maç daha")

        if kalici_hata_idler:
            print(f"\n[KALICI HATA] {len(kalici_hata_idler):,} maç 401/403 "
                  f"nedeniyle atlandı (checkpoint'e yazıldı):")
            for mid in list(kalici_hata_idler)[:20]:
                b = genis_veri.get(mid, {})
                print(f"  - {mid} | {b.get('Maç', '')} | "
                      f"{b.get('Arşiv Sayfa Durumu', '')}")
            if len(kalici_hata_idler) > 20:
                print(f"  ... ve {len(kalici_hata_idler) - 20} maç daha")

        basarisiz_kalan = toplam_mac - len(tamamlanan_idler)
        print(f"\n[SONUÇ] Tamamlanan (sayfa+oran): {len(tamamlanan_idler):,}/"
              f"{toplam_mac:,} | Başarısız kalan: {basarisiz_kalan:,} | "
              f"Kalıcı 401/403: {len(kalici_hata_idler):,}")

    except KeyboardInterrupt:
        print("\n[DURDURULDU] Mevcut veriler kaydedilecek...")
    except ErisimKisitlamaHatasi as hata:
        print(f"\n[ERİŞİM ENGELLENDİ] HTTP {hata.durum_kodu}: {hata}")
        if hata.durum_kodu in (401, 403):
            print("Sunucu erişimi reddetti. WARP/bağlantı durumunu kontrol edin; "
                  "bir süre sonra yeniden deneyin.")
        elif hata.durum_kodu == 429:
            print("İstek sınırı aşıldı. Otomatik bekleme/reset limiti dolmuş olabilir; "
                  "WARP bağlantısını kontrol edip soğuma süresinden sonra yeniden deneyin.")
        print("Mevcut ilerleme finally bölümünde kaydedilecek.")
    except Exception as hata:
        print(f"\n[İŞLEM HATASI] {hata}")
    finally:
        hatalar = hatalari_tekrarsiz_yap(hatalar)
        try:
            mac_log.kapat()
        except Exception:
            pass

        eski_liste = kayit.get("mac_listesi") if kayit else []
        if not df_mac.empty and df_mac_records:
            son_liste = df_mac_records
        elif eski_liste:
            son_liste = eski_liste
        else:
            son_liste = []

        try:
            json_ok, excel_ok, hata_ok = await guvenli_tam_kaydet(
                yollar, basla, son, genis_veri, hatalar, oran_kolonlari,
                mac_listesi_=son_liste,
                mac_listesi_durum_=mac_listesi_durum,
                excel_dosya_=excel_dosya, hata_dosya_=hata_dosya,
                mac_log_=mac_log, excel_yazilsin_=True,
                tamamlanan_gunler_=tamamlanan_gunler,
                kalici_hata_idler_=kalici_hata_idler,
            )
        except Exception as kayit_hatasi:
            print(f"[FINAL KAYIT HATASI] {kayit_hatasi}")
            json_ok = excel_ok = hata_ok = False

        if json_ok:
            print("[JSON] Ara kayıt diske yazıldı.")
            if not excel_ok:
                print("[UYARI] Excel yazılamadı (JSON korundu).")
            else:
                print(f"Excel: {excel_dosya.resolve()}")
            if not hata_ok:
                print("[UYARI] Hata raporu yazılamadı.")
            else:
                print(f"Hata raporu: {hata_dosya.resolve()}")
        else:
            print("[UYARI] Final JSON yazılamadı!")

        if (json_ok and not temp_silindi
                and ara_liste_dosyasi.exists()
                and _checkpoint_mac_listesi_dolu(yollar)):
            try:
                ara_liste_dosyasi.unlink()
                temp_silindi = True
                print("[MAÇ LİSTESİ] Geçici ara kayıt temizlendi.")
            except Exception as e:
                print(f"[UYARI] Geçici maç listesi silinemedi: {e}")
        elif ara_liste_dosyasi.exists() and not json_ok:
            print("[MAÇ LİSTESİ] Geçici ara kayıt KORUNDU "
                  "(checkpoint başarısız veya maç listesi doğrulanamadı).")

        print("-" * 76)
        print(f"Oturum süresi: {sure_formatla(time.time() - oturum_baslangici)}")
        tam_mac = sum(
            1 for mid in df_mac_ids
            if _tamamlanma_kriteri(genis_veri.get(mid, {}))
        )
        print(f"Tamamlanan (sayfa+oran): {tam_mac:,}/{toplam_mac:,}")
        if kalici_hata_idler:
            print(f"Kalıcı 401/403: {len(kalici_hata_idler):,}")
        print("=" * 76)


# ============================================================
# ÇALIŞTIR
# ============================================================

def _gruplari_coz(grup_str):
    if not grup_str or grup_str.strip().upper() in {"ALL", "HEPSI", "HEPSİ"}:
        return list(ULKE_GRUPLARI.keys()), "ALL"
    kodlar = [k.strip().upper() for k in grup_str.split(",") if k.strip()]
    gecersiz = [k for k in kodlar if k not in ULKE_GRUPLARI]
    if gecersiz:
        raise SystemExit(f"Geçersiz grup: {', '.join(gecersiz)}")
    gorulen = set()
    kodlar_temiz = []
    for k in kodlar:
        if k not in gorulen:
            gorulen.add(k)
            kodlar_temiz.append(k)
    return kodlar_temiz, "_".join(kodlar_temiz)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Maçkolik arşiv çekimi (v15.11)")
    parser.add_argument("--grup", "-g", default=None)
    parser.add_argument("--excel", action="store_true")
    args = parser.parse_args()

    AKTIF_GRUPLAR, AKTIF_ETIKET = _gruplari_coz(args.grup)
    AKTIF_LIGLER = set()
    for kod in AKTIF_GRUPLAR:
        AKTIF_LIGLER.update(ULKE_GRUPLARI[kod]["ligler"])

    print(f"[GRUP] {AKTIF_ETIKET} | "
          f"{', '.join(ULKE_GRUPLARI[k]['ad'] for k in AKTIF_GRUPLAR)} | "
          f"{len(AKTIF_LIGLER)} lig")

    try:
        baslangic, bitis = tarih_araligi_iste()
        if args.excel:
            jsondan_excel_uret(baslangic, bitis)
        else:
            asyncio.run(iddaa_bahis_oranlari_async(baslangic, bitis))
    except KeyboardInterrupt:
        print("\nİptal edildi. JSON korunur.")
    except Exception as hata:
        print(f"\nKRİTİK HATA: {hata}")
        raise